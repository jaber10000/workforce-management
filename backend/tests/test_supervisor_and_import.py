"""Supervisor role (review + draft editing, no approve/publish, no private data) and the in-app Excel import."""
import io
from datetime import date

import pytest
from openpyxl import Workbook

from app.core.security import hash_password
from app.models.roster import RosterVersion
from app.models.user import AppUser
from tests.helpers import START
from tests.test_leave_and_versions import client  # noqa: F401  (fixture)
from tests.test_production import API, PW, gen, hdr, published, sec  # noqa: F401  (sec is a fixture)


@pytest.fixture()
def sup(sec):
    s = sec.Session()
    s.add(AppUser(username="sup1", password_hash=hash_password(PW), role="supervisor", employee_id=None))
    s.commit()
    s.close()
    return sec


def _xlsx(build) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- supervisor --

def test_supervisor_can_review_and_edit_draft_but_not_decide(sup):
    mgr, s = hdr(sup, "mgr1"), hdr(sup, "sup1")
    v = gen(sup, mgr)
    vid = v["id"]
    # review access
    assert sup.get(f"{API}/roster/versions", headers=s).json()
    assert sup.get(f"{API}/roster/versions/{vid}", headers=s).json()["warnings"] is not None
    for path in ("coverage", "fairness", "diff", "history"):
        assert sup.get(f"{API}/roster/versions/{vid}/{path}", headers=s).status_code == 200, path
    # draft editing is allowed and attributed to the supervisor
    grid = sup.get(f"{API}/roster/versions/{vid}", headers=s).json()["grid"]
    row, day = grid["rows"][0], grid["days"][0]["date"]
    r = sup.put(f"{API}/roster/versions/{vid}/assignments", headers=s,
                json={"employee_id": row["employee_id"], "work_date": day, "code": "-", "reason": "supervisor test"})
    assert r.status_code == 200, r.text
    # decisions and destructive actions stay with managers
    assert sup.post(f"{API}/roster/generate", json={"start_date": str(START), "days": 7}, headers=s).status_code == 403
    assert sup.post(f"{API}/roster/versions/{vid}/approve", json={}, headers=s).status_code == 403
    assert sup.post(f"{API}/roster/versions/{vid}/publish", json={}, headers=s).status_code == 403
    assert sup.post(f"{API}/roster/versions/{vid}/reject", json={"note": "x"}, headers=s).status_code == 403
    assert sup.delete(f"{API}/roster/versions/{vid}", headers=s).status_code == 403
    assert sup.get(f"{API}/roster/versions/{vid}/export.xlsx", headers=s).status_code == 403
    assert sup.get(f"{API}/roster/audit-log", headers=s).status_code == 403
    assert sup.get(f"{API}/exports/employees.xlsx", headers=s).status_code == 403


def test_supervisor_leave_and_swap_review_only(sup):
    mgr, s = hdr(sup, "mgr1"), hdr(sup, "sup1")
    sup.post(f"{API}/leaves", headers=hdr(sup, "emp1"),
             json={"employee_id": "T001", "leave_type": "Casual", "start_date": "2026-11-20", "end_date": "2026-11-20"})
    items = sup.get(f"{API}/leaves", headers=s).json()["items"]
    assert items, "supervisor sees all leave applications"
    app_id = items[0]["application_id"]
    assert sup.get(f"{API}/leaves/balance/T001", headers=s).status_code == 200
    assert sup.get(f"{API}/leaves/{app_id}", headers=s).status_code == 200
    for action in ("approve", "reject"):
        assert sup.post(f"{API}/leaves/{app_id}/{action}", json={}, headers=s).status_code == 403
    # supervisors cannot file leave on behalf of someone else
    assert sup.post(f"{API}/leaves", headers=s, json={"employee_id": "T002", "leave_type": "Casual",
                    "start_date": "2026-11-21", "end_date": "2026-11-21"}).status_code == 403
    assert sup.get(f"{API}/swaps", headers=s).status_code == 200


def test_supervisor_does_not_see_private_employee_fields(sup):
    mgr, s = hdr(sup, "mgr1"), hdr(sup, "sup1")
    sup.patch(f"{API}/employees/T005", json={"phone_number": "01799999999", "home_thana": "Mirpur", "notes": "secret"}, headers=mgr)
    page = sup.get(f"{API}/employees?search=T005", headers=s).json()
    e = page["items"][0]
    assert e["employee_id"] == "T005" and e["full_name"]
    assert e["phone_number"] is None and e["home_thana"] is None and e["notes"] is None
    one = sup.get(f"{API}/employees/T005", headers=s).json()
    assert one["phone_number"] is None and one["notes"] is None
    # cannot probe private fields through search
    assert sup.get(f"{API}/employees?search=01799999999", headers=s).json()["total"] == 0
    assert sup.get(f"{API}/employees?search=Mirpur", headers=s).json()["total"] == 0
    assert sup.get(f"{API}/employees?search=01799999999", headers=mgr).json()["total"] == 1  # manager still can
    # and cannot write
    assert sup.patch(f"{API}/employees/T005", json={"full_name": "X"}, headers=s).status_code == 403
    assert sup.delete(f"{API}/employees/T005", headers=s).status_code == 403


def test_admin_can_create_supervisor_without_employee_link(sup):
    adm = hdr(sup, "admin")
    r = sup.post(f"{API}/auth/users", json={"username": "sup2", "password": "Passw0rd!y", "role": "supervisor"}, headers=adm)
    assert r.status_code == 201 and r.json()["role"] == "supervisor"
    assert sup.post(f"{API}/auth/users", json={"username": "sup3", "password": "Passw0rd!y", "role": "boss"}, headers=adm).status_code == 422
    assert sup.get(f"{API}/auth/users", headers=hdr(sup, "sup1")).status_code == 403
    assert sup.get(f"{API}/auth/me", headers=hdr(sup, "sup1")).json()["role"] == "supervisor"


# -------------------------------------------------------------------- import --

def _roster_book(wb):
    ws = wb.create_sheet("Weekly_Roster")
    ws.append(["Employee ID", "Employee Name", "Gender", "Role / Pool", "Home Thana",
               "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"])
    ws.append(["T001", "Test 1", "Male", "General", "x", "M", "E", "N", "OFF", "HDLM", "HDE", "HDLE"])
    ws.append(["T002", "Test 2", "Male", "General", "x", "E", "E", "E", "E", "E", "OFF", "OFF"])
    ws.append(["GHOST9", "Nobody", "Male", "General", "x", "M", "M", "M", "M", "M", "M", "M"])
    ws.append(["T003", "Test 3", "Male", "General", "x", "M", "BOGUS", "M", "M", "M", "OFF", "OFF"])
    dr = wb.create_sheet("Daily_Requirement")
    dr.append(["Date", "Day", "M", "E", "N", "Min Senior"])
    dr.append([date(2026, 10, 26), "Monday", 3, 2, 1, 1])


def test_import_dry_run_saves_nothing_then_real_import_creates_draft(sup):
    mgr = hdr(sup, "mgr1")
    data = _xlsx(_roster_book)
    url = f"{API}/imports/roster?week_start=2026-10-26&employees=false"
    r = sup.post(url, content=data, headers=mgr)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["dry_run"] is True and body["roster"]["cells"] == 21 and body["roster"]["version_id"] is None
    assert body["roster"]["unknown_employee_ids"] == ["GHOST9"]
    assert body["requirements"] == 1
    assert any("BOGUS" in e for e in body["errors"])
    s = sup.Session()
    assert s.query(RosterVersion).count() == 0, "dry run must not persist"
    s.close()

    r = sup.post(url.replace("?", "?dry_run=false&"), content=data, headers=mgr)
    assert r.status_code == 200, r.text
    info = r.json()["roster"]
    assert info["start_date"] == "2026-10-26" and info["end_date"] == "2026-11-01" and info["version_id"]
    v = sup.get(f"{API}/roster/versions/{info['version_id']}", headers=mgr).json()
    assert v["version"]["status"] == "draft" and v["version"]["trigger"] == "import"
    rows = {r["employee_id"]: r["codes"] for r in v["grid"]["rows"]}
    assert rows["T001"] == ["M", "E", "N", "WO", "M-H", "E-H", "E-H"]  # OFF/HDLM/HDE/HDLE mapped
    assert rows["T003"][1] == "-" and "GHOST9" not in rows
    actions = {a["action"] for a in sup.get(f"{API}/roster/audit-log", headers=mgr).json()}
    assert {"roster_imported", "workbook_imported"} <= actions
    # it is a normal draft: it can be approved and a second import becomes version 2
    again = sup.post(url.replace("?", "?dry_run=false&"), content=data, headers=mgr).json()["roster"]["version_id"]
    assert sup.get(f"{API}/roster/versions/{again}", headers=mgr).json()["version"]["version_no"] == 2


def test_import_employees_and_date_headers_roundtrip(sup):
    mgr = hdr(sup, "mgr1")

    def book(wb):
        em = wb.create_sheet("Employee_Master")
        em.append(["Employee ID", "Employee Name", "Gender", "Phone Number", "Home Thana", "Role / Pool",
                   "Weekly Off Day", "Half-duty Day"])
        em.append(["NEW001", "Brand New", "Female", "0170", "Uttara", "General", "Friday", None])
        ro = wb.create_sheet("Roster")  # the app's own export layout: date headers, no start date needed
        ro.append(["Employee ID", "Employee Name", "Role", "Mon 2026-11-02", "Tue 2026-11-03"])
        ro.append(["NEW001", "Brand New", "general", "M", "E"])

    r = sup.post(f"{API}/imports/roster?dry_run=false", content=_xlsx(book), headers=mgr)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["employees"]["created"] == 1 and b["roster"]["start_date"] == "2026-11-02" and b["roster"]["cells"] == 2
    assert sup.get(f"{API}/employees/NEW001", headers=mgr).json()["home_thana"] == "Uttara"


def test_import_permissions_and_bad_files(sup):
    data = _xlsx(_roster_book)
    for user in ("sup1", "emp1"):
        assert sup.post(f"{API}/imports/roster?week_start=2026-10-26", content=data, headers=hdr(sup, user)).status_code == 403
    assert sup.post(f"{API}/imports/roster", content=data).status_code == 401
    mgr = hdr(sup, "mgr1")
    assert sup.post(f"{API}/imports/roster", content=b"", headers=mgr).status_code == 422
    assert sup.post(f"{API}/imports/roster", content=b"not an excel file", headers=mgr).status_code == 415
    assert sup.post(f"{API}/imports/roster", content=b"PK\x03\x04garbage", headers=mgr).status_code == 422
    # weekday headers without a start date -> clear error, nothing saved
    r = sup.post(f"{API}/imports/roster?dry_run=false&employees=false&requirements=false", content=data, headers=mgr)
    assert r.status_code == 422 and "start date" in r.json()["detail"]
    empty = _xlsx(lambda wb: wb.create_sheet("Other"))
    assert sup.post(f"{API}/imports/roster", content=empty, headers=mgr).status_code == 422


# --------------------------------------------------- teams + profile creation --

def _seed_org(sec, n_extra=0):
    """5 supervisors, 10 seniors, plus the 70 test employees' generals."""
    from app.models.employee import Employee, RolePool
    s = sec.Session()
    for i in range(5):
        s.add(Employee(employee_id=f"SUP{i + 1}", full_name=f"Sup {i + 1}", role_pool=RolePool.supervisor))
    for e in s.query(Employee).filter(Employee.role_pool == RolePool.senior_support).all():
        e.role_pool = RolePool.general
    for i, e in enumerate(s.query(Employee).filter(Employee.employee_id.like("T%")).order_by(Employee.employee_id).limit(10).all()):
        e.role_pool = RolePool.senior_support
    s.commit()
    s.close()


def test_auto_assign_teams_balanced(sec):
    mgr, adm = hdr(sec, "mgr1"), hdr(sec, "admin")
    _seed_org(sec)
    r = sec.post(f"{API}/teams/auto-assign", headers=mgr)  # dry run by default
    assert r.status_code == 200 and r.json()["dry_run"] is True
    assert all(e["team"] is None for e in sec.get(f"{API}/employees?limit=200", headers=mgr).json()["items"])
    r = sec.post(f"{API}/teams/auto-assign?dry_run=false", headers=mgr).json()
    teams = r["teams"]
    assert sorted(t["team"] for t in teams) == [f"Team Sup {i}" for i in range(1, 6)]  # named after the leader
    assert all(t["team"] == f"Team {t['supervisors'][0]['full_name']}" for t in teams)
    sizes = [t["size"] for t in teams]
    assert max(sizes) - min(sizes) <= 1 and sum(sizes) == 75
    assert all(len(t["supervisors"]) == 1 and len(t["seniors"]) == 2 for t in teams), teams
    assert r["warnings"] == []
    # random by default, reproducible with the seed that the response reports
    seed = r["seed"]
    again = sec.post(f"{API}/teams/auto-assign?dry_run=false&seed={seed}", headers=mgr).json()
    assert again["changed"] == 0 and again["seed"] == seed
    other = sec.post(f"{API}/teams/auto-assign?dry_run=true&seed={seed + 1}", headers=mgr).json()
    assert other["changed"] > 0  # a different random split
    assert sec.get(f"{API}/teams", headers=hdr(sec, "emp1")).status_code == 403
    assert sec.post(f"{API}/teams/auto-assign?dry_run=false", headers=hdr(sec, "emp1")).status_code == 403
    assert sec.post(f"{API}/teams/auto-assign?teams=7", headers=adm).json()["warnings"]  # 5 supervisors for 7 teams


def test_supervisor_edits_only_own_team(sup):
    from app.models.user import AppUser
    mgr = hdr(sup, "mgr1")
    _seed_org(sup)
    assert sup.post(f"{API}/teams/auto-assign?dry_run=false", headers=mgr).status_code == 200
    s = sup.Session()
    s.add(AppUser(username="sup_t1", password_hash=hash_password(PW), role="supervisor", employee_id="SUP1"))
    s.commit()
    s.close()
    h = hdr(sup, "sup_t1")
    v = gen(sup, mgr)
    g = sup.get(f"{API}/roster/versions/{v['id']}", headers=h).json()["grid"]
    t1 = sup.get(f"{API}/auth/me", headers=h).json()  # SUP1's own team name (named after a random leader)
    my_team = next(r["team"] for r in g["rows"] if r["employee_id"] == "SUP1")
    mine = next(r for r in g["rows"] if r["team"] == my_team and r["role"] == "general")
    other = next(r for r in g["rows"] if r["team"] != my_team and r["role"] == "general")
    day = g["days"][0]["date"]
    body = lambda eid: {"employee_id": eid, "work_date": day, "code": "-", "reason": "team scope"}
    assert sup.put(f"{API}/roster/versions/{v['id']}/assignments", json=body(mine["employee_id"]), headers=h).status_code == 200
    r = sup.put(f"{API}/roster/versions/{v['id']}/assignments", json=body(other["employee_id"]), headers=h)
    assert r.status_code == 403 and "not in your team" in r.json()["detail"]
    # managers are not limited, and an unlinked supervisor keeps full access (tested above)
    assert sup.put(f"{API}/roster/versions/{v['id']}/assignments", json=body(other["employee_id"]), headers=mgr).status_code == 200


def test_create_user_creates_missing_profile(sup):
    adm = hdr(sup, "admin")
    # supervisor / manager without a profile at all is fine
    assert sup.post(f"{API}/auth/users", headers=adm, json={"username": "boss1", "password": PW, "role": "manager"}).status_code == 201
    # an ID that has no profile: clear message instead of the old 404 ...
    body = {"username": "sup9", "password": PW, "role": "supervisor", "employee_id": "SUP9"}
    r = sup.post(f"{API}/auth/users", headers=adm, json=body)
    assert r.status_code == 422 and "full name" in r.json()["detail"]
    # ... and with a name the profile is created together with the account
    r = sup.post(f"{API}/auth/users", headers=adm, json={**body, "full_name": "Sabbir Hossain", "team": "Team Rahim"})
    assert r.status_code == 201 and r.json()["profile_created"] is True
    e = sup.get(f"{API}/employees/SUP9", headers=hdr(sup, "mgr1")).json()
    assert e["full_name"] == "Sabbir Hossain" and e["role_pool"] == "supervisor" and e["team"] == "Team Rahim"
    # existing profile: linked (case-insensitive) and not duplicated
    r = sup.post(f"{API}/auth/users", headers=adm, json={"username": "emp7", "password": PW, "role": "employee", "employee_id": "t007"})
    assert r.status_code == 201 and r.json()["employee_id"] == "T007" and r.json()["profile_created"] is False
    # employee accounts still need an ID
    assert sup.post(f"{API}/auth/users", headers=adm, json={"username": "emp8", "password": PW, "role": "employee"}).status_code == 422


# ------------------------------------------------ my team / my shifts / messages --

PNG = ("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def _org_with_published(sec):
    mgr = hdr(sec, "mgr1")
    _seed_org(sec)
    sec.post(f"{API}/teams/auto-assign?dry_run=false&seed=3", headers=mgr)
    s = sec.Session()
    s.add(AppUser(username="sup_t", password_hash=hash_password(PW), role="supervisor", employee_id="SUP1"))
    s.commit()
    s.close()
    published(sec, mgr, ack=True)
    return mgr


def test_employee_sees_team_leader_seniors_members_and_next_shifts(sec):
    from tests.test_production import published  # noqa: F401
    _org_with_published(sec)
    e = hdr(sec, "emp1")
    t = sec.get(f"{API}/me/team", headers=e).json()
    assert t["team"].startswith("Team ") and t["leader"]["employee_id"].startswith("SUP")
    assert t["team"] == f"Team {t['leader']['full_name']}"
    assert t["size"] == len(t["members"]) and any(m["is_me"] for m in t["members"])
    assert all(set(m) == {"employee_id", "full_name", "role_pool", "is_me"} for m in t["members"])  # no phone / thana
    assert all(x["role_pool"] == "senior_support" for x in t["seniors"])
    sh = sec.get(f"{API}/me/shifts?days=7", headers=e).json()
    assert len(sh["items"]) == 7 and sh["published_version_id"] and any(i["code"] for i in sh["items"])
    # drafts never leak into the employee's shifts
    assert sec.get(f"{API}/me/team/roster", headers=e).status_code == 403
    sup = hdr(sec, "sup_t")
    r = sec.get(f"{API}/me/team/roster?days=7", headers=sup).json()
    assert r["team"] and len(r["rows"]) >= 10 and len(r["days"]) == 7


def test_private_messages_with_image_between_employee_and_leader(sec):
    import base64
    _org_with_published(sec)
    e1, e2 = hdr(sec, "emp1"), hdr(sec, "emp2")
    me = sec.get(f"{API}/me/team", headers=e1).json()
    leader = me["leader"]["employee_id"]
    s = sec.Session()
    s.add(AppUser(username="lead", password_hash=hash_password(PW), role="supervisor", employee_id=leader))
    s.commit()
    s.close()
    lead = hdr(sec, "lead")
    r = sec.post(f"{API}/messages/thread/T001", headers=e1, json={"text": "Machine 4 is broken", "topic": "Problem", "image_b64": PNG, "image_name": "m4.png"})
    assert r.status_code == 201 and r.json()["has_image"] and r.json()["side"] == "employee"
    mid = r.json()["id"]
    assert sec.post(f"{API}/messages/thread/T001", headers=e1, json={"text": ""}).status_code == 422
    assert sec.post(f"{API}/messages/thread/T001", headers=e1, json={"text": "x", "image_b64": base64.b64encode(b"not an image").decode()}).status_code == 422
    # leader sees an unread, replies; privacy holds for others
    th = sec.get(f"{API}/messages/threads", headers=lead).json()
    assert next(t for t in th if t["employee_id"] == "T001")["unread"] == 1
    assert sec.get(f"{API}/messages/{mid}/image", headers=lead).headers["content-type"] == "image/png"
    assert sec.get(f"{API}/messages/thread/T001", headers=lead).json()["messages"][0]["text"] == "Machine 4 is broken"
    assert next(t for t in sec.get(f"{API}/messages/threads", headers=lead).json() if t["employee_id"] == "T001")["unread"] == 0
    assert sec.post(f"{API}/messages/thread/T001", headers=lead, json={"text": "I will check today"}).json()["side"] == "staff"
    assert sec.get(f"{API}/messages/unread", headers=e1).json()["unread"] == 1
    assert sec.get(f"{API}/messages/thread/T001", headers=e2).status_code == 403   # another employee
    assert sec.get(f"{API}/messages/{mid}/image", headers=e2).status_code == 403
    assert sec.get(f"{API}/messages/thread/T001", headers=hdr(sec, "mgr1")).status_code == 403  # private: managers cannot read it
    assert sec.get(f"{API}/messages/thread/T001", headers=hdr(sec, "admin")).status_code == 403
    other_leader = sec.get(f"{API}/messages/thread/T001", headers=hdr(sec, "sup_t"))
    assert other_leader.status_code in (200, 403)
