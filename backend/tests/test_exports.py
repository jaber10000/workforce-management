"""Step 7: Excel exports - content, formatting, validation, version/date metadata, permissions, import round trip."""
import io
from datetime import date, datetime, timezone

from openpyxl import load_workbook

from app.scripts.import_employees import import_workbook
from app.services import export_service as ex
from tests.helpers import START
from tests.test_leave_and_versions import client  # noqa: F401  (fixture)
from tests.test_production import API, gen, grid, hdr, published, sec  # noqa: F401  (fixtures)

WEEKDAY_LABELS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def wb_of(resp):
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(resp.content))


def info(wb):
    return {r[0].value: r[1].value for r in wb["Info"].iter_rows() if r[0].value}


def test_weekly_roster_export_content_version_and_date(sec):
    mgr, emp = hdr(sec, "mgr1"), hdr(sec, "emp1")
    v = published(sec, mgr)
    r = sec.get(f"{API}/exports/roster/weekly.xlsx?week_start={START}", headers=emp)  # employees may export published
    wb = wb_of(r)
    meta = info(wb)
    assert meta["Roster version"] == f"v{v['version_no']} (id {v['id']})"
    assert meta["Version status"].startswith("OFFICIAL")
    assert meta["Exported on (date)"].date() == datetime.now(timezone.utc).date()
    assert meta["Exported by"] == "emp1" and meta["Week"] == f"{START} to {date.fromordinal(START.toordinal() + 6)}"
    disp = r.headers["content-disposition"]
    assert f"v{v['version_no']}_published" in disp and "exported-" in disp and str(START) in disp

    ws = wb["Weekly Roster"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ex.WEEKLY_FIXED and len(header) == 3 + 7
    assert header[3] == f"{WEEKDAY_LABELS[START.weekday()]} {START:%d-%b}"
    api_grid = grid(sec, mgr, v["id"])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=10, values_only=True):
        if row[0] in api_grid:
            assert list(row[3:10]) == api_grid[row[0]][:7]  # identical to the roster API
            assert all(code in ex.ROSTER_CODES for code in row[3:10])
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and len(ws.data_validations.dataValidation) == 1
    labels = [ws.cell(row=ws.max_row - 2 + i, column=1).value for i in range(3)]
    assert labels == ["Morning", "Evening", "Night"]  # on-duty totals block


def test_weekly_export_draft_is_marked_unofficial_and_hidden_from_employees(sec):
    mgr, emp = hdr(sec, "mgr1"), hdr(sec, "emp1")
    draft = gen(sec, mgr)
    r = sec.get(f"{API}/exports/roster/weekly.xlsx?version_id={draft['id']}&week_start={START}", headers=mgr)
    assert info(wb_of(r))["Version status"] == "NOT OFFICIAL - DRAFT" and "_draft_" in r.headers["content-disposition"]
    assert sec.get(f"{API}/exports/roster/weekly.xlsx?version_id={draft['id']}&week_start={START}", headers=emp).status_code == 403
    assert sec.get(f"{API}/exports/roster/weekly.xlsx?week_start={START}", headers=emp).status_code == 404  # nothing published
    assert sec.get(f"{API}/exports/roster/weekly.xlsx").status_code == 401


def test_weekly_export_rejects_week_outside_roster(sec):
    mgr = hdr(sec, "mgr1")
    v = gen(sec, mgr)
    r = sec.get(f"{API}/exports/roster/weekly.xlsx?version_id={v['id']}&week_start=2027-03-01", headers=mgr)
    assert r.status_code == 422 and "covers" in r.json()["detail"]


def test_employee_list_export_columns_validation_and_formula_safety(sec):
    mgr = hdr(sec, "mgr1")
    evil = {"employee_id": "X1", "full_name": '=HYPERLINK("http://evil.example","click")', "phone_number": "01712345678",
            "role_pool": "supervisor", "is_shift_eligible_night": True}
    assert sec.post(f"{API}/employees?auto_roster=false", json=evil, headers=mgr).status_code == 201
    assert sec.post(f"{API}/employees?auto_roster=false", json={"employee_id": "GONE", "full_name": "Left"}, headers=mgr).status_code == 201
    assert sec.delete(f"{API}/employees/GONE?auto_roster=false", headers=mgr).status_code == 204

    wb = wb_of(sec.get(f"{API}/exports/employees.xlsx", headers=mgr))
    ws = wb["Employees"]
    assert ex.check_columns(ws, ex.EXPECTED_COLUMNS["employees"]) == []
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    assert "GONE" not in rows and "T001" in rows  # inactive excluded by default
    x = rows["X1"]
    assert x[1].value.startswith("'=") and x[3].value == "01712345678" and x[3].number_format == "@"
    assert x[5].value == "supervisor" and x[11].value == "Yes"
    assert len(ws.data_validations.dataValidation) >= 6  # gender, role, weekdays, yes/no columns
    assert ws.auto_filter.ref.startswith("A1:") and ws.freeze_panes == "A2"
    assert info(wb)["Employees"] == ws.max_row - 1

    inactive = wb_of(sec.get(f"{API}/exports/employees.xlsx?is_active=false", headers=mgr))["Employees"]
    assert [r[0].value for r in inactive.iter_rows(min_row=2)] == ["GONE"]
    assert sec.get(f"{API}/exports/employees.xlsx?role_pool=nonsense", headers=mgr).status_code == 422


def test_exported_employee_file_can_be_imported_back(sec, sqlite_session):
    mgr = hdr(sec, "mgr1")
    sec.post(f"{API}/employees?auto_roster=false", headers=mgr,
             json={"employee_id": "X1", "full_name": "=SUM(1)", "phone_number": "0171", "weekly_off_day": "Friday"})
    r = sec.get(f"{API}/exports/employees.xlsx", headers=mgr)
    path = io.BytesIO(r.content)
    created, updated, skipped, errors = import_workbook(path, db=sqlite_session)
    assert errors == [] and created == 71 and updated == 0
    from app.models.employee import Employee
    x = sqlite_session.query(Employee).filter_by(employee_id="X1").one()
    assert x.full_name == "=SUM(1)" and x.phone_number == "0171" and x.weekly_off_day.value == "Friday"  # guard undone
    t1 = sqlite_session.query(Employee).filter_by(employee_id="T001").one()
    assert t1.role_pool.value == "senior_support" and t1.is_shift_eligible_night is True
    again = import_workbook(io.BytesIO(r.content), db=sqlite_session)
    assert again[0] == 0 and again[1] == 71  # idempotent: updates, no duplicates


def test_import_reports_duplicates_and_bad_headers(tmp_path, sqlite_session):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Employee_Master"
    ws.append(["Employee ID", "Employee Name"])
    ws.append(["A1", "One"]); ws.append(["a1", "One again"]); ws.append(["A2", "Two"]); ws.append([None, "No id"])
    p = tmp_path / "e.xlsx"
    wb.save(p)
    created, updated, skipped, errors = import_workbook(p, db=sqlite_session)
    assert (created, skipped) == (2, 1) and len(errors) == 1 and "duplicate Employee ID" in errors[0]
    wb2 = Workbook(); wb2.active.title = "Employee_Master"; wb2.active.append(["Name only"])
    wb2.save(tmp_path / "bad.xlsx")
    import pytest
    with pytest.raises(ValueError, match="Missing required column"):
        import_workbook(tmp_path / "bad.xlsx", db=sqlite_session)


def test_leave_report_export_filters_totals_and_permissions(sec):
    mgr, emp1, emp2 = hdr(sec, "mgr1"), hdr(sec, "emp1"), hdr(sec, "emp2")
    a = sec.post(f"{API}/leaves", headers=emp1, json={"employee_id": "T001", "leave_type": "Casual",
                                                      "start_date": "2026-11-05", "end_date": "2026-11-06"}).json()
    sec.post(f"{API}/leaves", headers=emp2, json={"employee_id": "T002", "leave_type": "Sick",
                                                  "start_date": "2026-11-10", "end_date": "2026-11-10", "reason": "=1+1"})
    sec.post(f"{API}/leaves/{a['application_id']}/approve", json={"auto_roster": False}, headers=mgr)

    wb = wb_of(sec.get(f"{API}/exports/leaves.xlsx", headers=mgr))
    ws = wb["Leave Applications"]
    assert ex.check_columns(ws, ex.EXPECTED_COLUMNS["leaves"]) == []
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == ["LV0001", "LV0002"] and rows[0][4].date() == date(2026, 11, 5) and rows[0][6] == 2
    assert rows[1][9] == "'=1+1" and rows[0][8] == "Approved" and rows[1][8] == "Pending"
    assert info(wb)["Applications"] == 2
    summ = wb["Summary"]
    assert summ["E2"].value == 2 and summ["F3"].value == 1  # approved days / pending days

    only = wb_of(sec.get(f"{API}/exports/leaves.xlsx?status=Approved&employee_id=T001", headers=mgr))
    assert only["Leave Applications"].max_row == 2 and info(only)["Filter: status"] == "Approved"
    rng = wb_of(sec.get(f"{API}/exports/leaves.xlsx?from_date=2026-11-08&to_date=2026-11-30", headers=mgr))
    assert rng["Leave Applications"].max_row == 2  # only the 10 Nov leave overlaps
    assert sec.get(f"{API}/exports/leaves.xlsx?status=Bogus", headers=mgr).status_code == 422
    assert sec.get(f"{API}/exports/leaves.xlsx?from_date=2026-12-01&to_date=2026-11-01", headers=mgr).status_code == 422
    assert sec.get(f"{API}/exports/leaves.xlsx", headers=emp1).status_code == 403  # employees: no leave report


def test_coverage_export_and_column_contract(sec):
    mgr, emp = hdr(sec, "mgr1"), hdr(sec, "emp1")
    v = published(sec, mgr)
    r = sec.get(f"{API}/exports/coverage.xlsx?version_id={v['id']}", headers=mgr)
    wb = wb_of(r)
    ws = wb["Shift Coverage"]
    assert ex.check_columns(ws, ex.EXPECTED_COLUMNS["coverage"]) == []
    data = [row for row in ws.iter_rows(min_row=2, max_row=15, values_only=True)]
    assert len(data) == 14 and data[0][0].date() == START and data[0][1] == "Monday"
    for row in data:  # shortage columns are arithmetically consistent with required/assigned
        assert row[4] == max(0, row[2] - row[3]) and row[7] == max(0, row[5] - row[6]) and row[10] == max(0, row[8] - row[9])
    assert len(ws.conditional_formatting) >= 4
    meta = info(wb)
    assert meta["Roster version"].startswith("v") and "Days with shortage" in meta and meta["Exported by"] == "mgr1"
    assert sec.get(f"{API}/exports/coverage.xlsx", headers=emp).status_code == 403

    broken = load_workbook(io.BytesIO(r.content))["Shift Coverage"]
    broken["C1"].value = "Morning Needed"
    problems = ex.check_columns(broken, ex.EXPECTED_COLUMNS["coverage"])
    assert "missing column 'Morning Required'" in problems and "unexpected column 'Morning Needed'" in problems


def test_every_export_is_audited(sec):
    mgr = hdr(sec, "mgr1")
    sec.get(f"{API}/exports/employees.xlsx", headers=mgr)
    log = sec.get(f"{API}/roster/audit-log?action=report_exported", headers=mgr).json()
    assert log and log[0]["actor"] == "mgr1" and log[0]["entity_id"] == "employee_list"
    assert log[0]["details"]["file"].startswith("employee_list_exported-")
