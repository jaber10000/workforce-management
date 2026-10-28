"""Approval workflow, permissions, shift swap, audit history and Excel export (auth ON)."""
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.api.routes import auth as auth_routes
from app.core.config import settings
from app.core.security import hash_password
from app.models.user import AppUser
from tests.helpers import START
from tests.test_leave_and_versions import Q, TODAY, client, cells  # noqa: F401  (client is a fixture)

API = "/api/v1"
PW = "Passw0rd!x"


@pytest.fixture()
def sec(client, monkeypatch):
    monkeypatch.setattr(settings, "auth_required", True)
    monkeypatch.setattr(settings, "require_approval", True)
    auth_routes._failures.clear()
    s = client.Session()
    for name, role, emp in [("admin", "admin", None), ("mgr1", "manager", None), ("mgr2", "manager", None),
                            ("emp1", "employee", "T001"), ("emp2", "employee", "T002")]:
        s.add(AppUser(username=name, password_hash=hash_password(PW), role=role, employee_id=emp))
    s.commit()
    s.close()
    return client


def hdr(c, user, pw=PW):
    r = c.post(f"{API}/auth/login", json={"username": user, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def gen(c, h, **extra):
    r = c.post(f"{API}/roster/generate", json={"start_date": str(START), "days": 14, **Q, **extra}, headers=h)
    assert r.status_code == 201, r.text
    return r.json()


def grid(c, h, vid):
    g = c.get(f"{API}/roster/versions/{vid}", headers=h).json()["grid"]
    return {r["employee_id"]: r["codes"] for r in g["rows"]}


def published(c, h, ack=False):
    v = gen(c, h)
    assert c.post(f"{API}/roster/versions/{v['id']}/approve", json={"acknowledge_shortage": ack}, headers=h).status_code == 200
    r = c.post(f"{API}/roster/versions/{v['id']}/publish", json={"acknowledge_shortage": ack}, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


# ------------------------------------------------------------------- auth --

def test_login_required_wrong_password_and_lockout(sec):
    assert sec.get(f"{API}/employees").status_code == 401
    assert sec.get(f"{API}/employees", headers={"Authorization": "Bearer garbage"}).status_code == 401
    assert sec.post(f"{API}/auth/login", json={"username": "mgr1", "password": "nope"}).status_code == 401
    for _ in range(settings.login_max_failures):
        sec.post(f"{API}/auth/login", json={"username": "emp2", "password": "bad"})
    assert sec.post(f"{API}/auth/login", json={"username": "emp2", "password": PW}).status_code == 429
    h = hdr(sec, "mgr1")
    assert sec.get(f"{API}/auth/me", headers=h).json()["role"] == "manager"
    assert sec.get("/health").status_code == 200 and sec.get("/ready").json() == {"status": "ready"}


def test_disabled_user_and_user_admin_rules(sec):
    admin, mgr = hdr(sec, "admin"), hdr(sec, "mgr1")
    body = {"username": "newmgr", "password": PW, "role": "manager"}
    assert sec.post(f"{API}/auth/users", json=body, headers=mgr).status_code == 403  # only admin
    assert sec.post(f"{API}/auth/users", json=body, headers=admin).status_code == 201
    assert sec.post(f"{API}/auth/users", json=body, headers=admin).status_code == 409
    assert sec.post(f"{API}/auth/users", json={**body, "username": "e9", "role": "employee"}, headers=admin).status_code == 422
    assert sec.post(f"{API}/auth/users", json={**body, "username": "short", "password": "abc"}, headers=admin).status_code == 422
    tok = hdr(sec, "newmgr")
    assert sec.patch(f"{API}/auth/users/newmgr", json={"is_active": False}, headers=admin).status_code == 200
    assert sec.get(f"{API}/employees", headers=tok).status_code == 401  # effective immediately, token or not
    assert sec.patch(f"{API}/auth/users/admin", json={"role": "employee"}, headers=admin).status_code == 409
    emp = hdr(sec, "emp1")
    assert sec.post(f"{API}/auth/change-password", json={"current_password": "x", "new_password": "Another1!pw"}, headers=emp).status_code == 401
    assert sec.post(f"{API}/auth/change-password", json={"current_password": PW, "new_password": "Another1!pw"}, headers=emp).status_code == 204
    assert sec.post(f"{API}/auth/login", json={"username": "emp1", "password": PW}).status_code == 401
    assert sec.post(f"{API}/auth/login", json={"username": "emp1", "password": "Another1!pw"}).status_code == 200


def test_role_permissions(sec):
    mgr, emp, emp2 = hdr(sec, "mgr1"), hdr(sec, "emp1"), hdr(sec, "emp2")
    # employee: read-only directory, nothing that changes the roster
    assert sec.get(f"{API}/employees", headers=emp).status_code == 200
    new = {"employee_id": "NEW9", "full_name": "X"}
    assert sec.post(f"{API}/employees?auto_roster=false", json=new, headers=emp).status_code == 403
    assert sec.post(f"{API}/employees?auto_roster=false", json=new, headers=mgr).status_code == 201
    assert sec.post(f"{API}/roster/generate", json={"start_date": str(START), **Q}, headers=emp).status_code == 403
    assert sec.get(f"{API}/roster/audit-log", headers=emp).status_code == 403
    assert sec.put(f"{API}/roster/requirements", json=[], headers=emp).status_code == 403
    # draft is invisible to employees, published is visible
    v = gen(sec, mgr)
    assert sec.get(f"{API}/roster/versions", headers=emp).json() == []
    assert sec.get(f"{API}/roster/versions/{v['id']}", headers=emp).status_code == 403
    sec.post(f"{API}/roster/versions/{v['id']}/approve", headers=mgr)
    sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr)
    assert [x["id"] for x in sec.get(f"{API}/roster/versions", headers=emp).json()] == [v["id"]]
    assert sec.get(f"{API}/roster/versions/{v['id']}", headers=emp).status_code == 200
    assert sec.get(f"{API}/roster/versions/{v['id']}/export.xlsx", headers=emp).status_code == 403
    # leave: own only
    mine = {"employee_id": "T001", "leave_type": "Casual", "start_date": "2026-11-05", "end_date": "2026-11-05"}
    assert sec.post(f"{API}/leaves", json={**mine, "employee_id": "T002"}, headers=emp).status_code == 403
    lid = sec.post(f"{API}/leaves", json=mine, headers=emp).json()["application_id"]
    assert sec.get(f"{API}/leaves/{lid}", headers=emp2).status_code == 403
    assert [x["application_id"] for x in sec.get(f"{API}/leaves", headers=emp).json()["items"]] == [lid]
    assert sec.get(f"{API}/leaves", headers=emp2).json()["items"] == []
    assert sec.get(f"{API}/leaves/balance/T001", headers=emp2).status_code == 403
    assert sec.post(f"{API}/leaves/{lid}/approve", headers=emp).status_code == 403
    assert sec.post(f"{API}/leaves/{lid}/preview", headers=emp).status_code == 403
    r = sec.post(f"{API}/leaves/{lid}/approve", json={"decided_by": "someone-else", "auto_roster": False}, headers=mgr)
    assert r.json()["leave"]["decided_by"] == "mgr1"  # actor comes from the login, not the request body
    assert sec.post(f"{API}/leaves/{lid}/cancel", headers=emp).status_code == 403  # approved: manager only


# ------------------------------------------------- approval workflow --

def test_draft_review_approve_publish_flow(sec):
    mgr, mgr2 = hdr(sec, "mgr1"), hdr(sec, "mgr2")
    v = gen(sec, mgr)
    assert v["status"] == "draft" and v["created_by"] == "mgr1"
    r = sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr)
    assert r.status_code == 409 and "Approve the draft first" in r.json()["detail"]
    ok = sec.post(f"{API}/roster/versions/{v['id']}/approve", json={"note": "looks good"}, headers=mgr).json()
    assert ok["status"] == "approved" and ok["approved_by"] == "mgr1" and ok["review_note"] == "looks good"
    assert sec.post(f"{API}/roster/versions/{v['id']}/approve", headers=mgr).status_code == 409  # already approved
    edit = {"employee_id": "T001", "work_date": "2026-11-03", "code": "R", "reason": "test change"}
    assert sec.put(f"{API}/roster/versions/{v['id']}/assignments", json=edit, headers=mgr).status_code == 409  # locked
    assert sec.delete(f"{API}/roster/versions/{v['id']}", headers=mgr).status_code == 409
    pub = sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr2).json()
    assert pub["status"] == "published" and pub["published_by"] == "mgr2" and pub["approved_by"] == "mgr1"

    # a newer published version supersedes the old one; history stays readable
    v2 = gen(sec, mgr)
    sec.post(f"{API}/roster/versions/{v2['id']}/approve", headers=mgr)
    sec.post(f"{API}/roster/versions/{v2['id']}/publish", headers=mgr)
    statuses = {x["id"]: x["status"] for x in sec.get(f"{API}/roster/versions", headers=mgr).json()}
    assert statuses == {v["id"]: "superseded", v2["id"]: "published"}
    assert sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr).status_code == 409


def test_reject_needs_note_and_blocks_publish(sec):
    mgr = hdr(sec, "mgr1")
    v = gen(sec, mgr)
    assert sec.post(f"{API}/roster/versions/{v['id']}/reject", json={"note": "x"}, headers=mgr).status_code == 422
    r = sec.post(f"{API}/roster/versions/{v['id']}/reject", json={"note": "night cover too thin"}, headers=mgr).json()
    assert r["status"] == "rejected" and r["review_note"] == "night cover too thin"
    assert sec.post(f"{API}/roster/versions/{v['id']}/approve", headers=mgr).status_code == 409
    assert sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr).status_code == 409
    assert sec.delete(f"{API}/roster/versions/{v['id']}", headers=mgr).status_code == 204  # rejected drafts may be removed


def test_shortage_needs_acknowledgement_at_approval(sec):
    mgr = hdr(sec, "mgr1")
    reqs = [{"work_date": str(START + timedelta(days=k)), "morning_required": 14, "evening_required": 10,
             "night_required": 30, "min_senior": 1} for k in range(7)]
    v = sec.post(f"{API}/roster/generate", json={"start_date": str(START), "days": 7, **Q, "requirements": reqs}, headers=mgr).json()
    r = sec.post(f"{API}/roster/versions/{v['id']}/approve", headers=mgr)
    assert r.status_code == 409 and "shortage" in r.json()["detail"]
    assert sec.post(f"{API}/roster/versions/{v['id']}/approve", json={"acknowledge_shortage": True}, headers=mgr).status_code == 200
    assert sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr).status_code == 409  # still must acknowledge
    assert sec.post(f"{API}/roster/versions/{v['id']}/publish", json={"acknowledge_shortage": True}, headers=mgr).status_code == 200


def test_four_eyes_rule(sec, monkeypatch):
    monkeypatch.setattr(settings, "separate_approver", True)
    mgr, mgr2 = hdr(sec, "mgr1"), hdr(sec, "mgr2")
    v = gen(sec, mgr)
    sec.post(f"{API}/roster/versions/{v['id']}/approve", headers=mgr)
    assert sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr).status_code == 403
    assert sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr2).status_code == 200


# -------------------------------------------------------------- swaps --

def find_pair(sec, mgr, vid, d):
    """Probe (as manager) for two people who may legally swap on date d; returns (a, b) and cancels the probe."""
    codes = grid(sec, mgr, vid)
    i = (d - START).days
    ms = [e for e, c in codes.items() if c[i] == "M"]
    es = [e for e, c in codes.items() if c[i] == "E"]
    for a in ms[:12]:
        for b in es[:12]:
            r = sec.post(f"{API}/swaps", json={"requester_id": a, "counterpart_id": b, "work_date": str(d)}, headers=mgr)
            if r.status_code == 201:
                sec.post(f"{API}/swaps/{r.json()['swap_id']}/cancel", headers=mgr)
                return a, b
    pytest.fail("no swappable pair found")


def test_swap_full_flow_creates_new_published_version(sec):
    mgr = hdr(sec, "mgr1")
    v = published(sec, mgr)
    d = date(2026, 11, 4)
    a, b = find_pair(sec, mgr, v["id"], d)
    s = sec.Session()
    for name, emp in (("swa", a), ("swb", b)):
        s.add(AppUser(username=name, password_hash=hash_password(PW), role="employee", employee_id=emp))
    s.commit(); s.close()
    ha, hb = hdr(sec, "swa"), hdr(sec, "swb")
    before = grid(sec, mgr, v["id"])
    i = (d - START).days

    # employees can only act for themselves
    assert sec.post(f"{API}/swaps", json={"requester_id": b, "counterpart_id": a, "work_date": str(d)}, headers=ha).status_code == 403
    assert sec.post(f"{API}/swaps", json={"counterpart_id": b, "work_date": "2026-10-28"}, headers=ha).status_code == 409  # past
    assert sec.post(f"{API}/swaps", json={"counterpart_id": a, "work_date": str(d)}, headers=ha).status_code == 422  # self
    r = sec.post(f"{API}/swaps", json={"counterpart_id": b, "work_date": str(d), "reason": "family event"}, headers=ha)
    assert r.status_code == 201 and r.json()["status"] == "pending_counterpart" and r.json()["requester_id"] == a
    sid = r.json()["swap_id"]
    assert sec.post(f"{API}/swaps", json={"counterpart_id": b, "work_date": str(d)}, headers=ha).status_code == 409  # duplicate
    assert sec.post(f"{API}/swaps/{sid}/approve", headers=ha).status_code == 403            # not a manager
    assert sec.post(f"{API}/swaps/{sid}/accept", headers=ha).status_code == 403             # requester cannot accept
    assert sec.post(f"{API}/swaps/{sid}/approve", headers=mgr).status_code == 409           # counterpart has not agreed
    assert sec.post(f"{API}/swaps/{sid}/accept", headers=hb).json()["status"] == "pending_manager"
    res = sec.post(f"{API}/swaps/{sid}/approve", json={"note": "ok"}, headers=mgr)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["swap"]["status"] == "approved" and body["swap"]["decided_by"] == "mgr1"

    new_id = body["new_version_id"]
    after = grid(sec, mgr, new_id)
    assert after[a][i] == before[b][i] and after[b][i] == before[a][i]       # shifts exchanged
    changed = [(e, k) for e in before for k in range(14) if before[e][k] != after[e][k]]
    assert sorted(changed) == sorted([(a, i), (b, i)])                       # nothing else moved
    vs = {x["id"]: x for x in sec.get(f"{API}/roster/versions", headers=mgr).json()}
    assert vs[v["id"]]["status"] == "superseded" and vs[new_id]["status"] == "published"
    assert vs[new_id]["trigger"] == "shift_swap" and vs[new_id]["base_version_id"] == v["id"]
    assert grid(sec, mgr, v["id"]) == before                                  # old version intact
    # approved swaps cannot be re-decided; employees see their own swaps only
    assert sec.post(f"{API}/swaps/{sid}/reject", headers=mgr).status_code == 409
    mine = sec.get(f"{API}/swaps", headers=ha).json()
    assert sid in [x["swap_id"] for x in mine] and all(a in (x["requester_id"], x["counterpart_id"]) for x in mine)
    other = sec.get(f"{API}/swaps", headers=hdr(sec, "emp2")).json()  # T002 sees only swaps they are part of
    assert all("T002" in (x["requester_id"], x["counterpart_id"]) for x in other)
    log = sec.get(f"{API}/roster/audit-log", params={"entity_id": sid}, headers=mgr).json()
    assert [x["action"] for x in log][::-1] == ["swap_requested", "swap_accepted", "swap_approved"]


def test_swap_rules_are_enforced_at_approval(sec):
    mgr = hdr(sec, "mgr1")
    v = published(sec, mgr)
    d = date(2026, 11, 4)
    codes = grid(sec, mgr, v["id"])
    i = (d - START).days
    ms = [e for e, c in codes.items() if c[i] == "M"]
    off = next(e for e, c in codes.items() if c[i] == "WO")
    # same shift / not on duty / no published roster / half vs full duty are refused with a reason
    r = sec.post(f"{API}/swaps", json={"requester_id": ms[0], "counterpart_id": ms[1], "work_date": str(d)}, headers=mgr)
    assert r.status_code == 409 and "same shift" in r.json()["detail"]
    r = sec.post(f"{API}/swaps", json={"requester_id": ms[0], "counterpart_id": off, "work_date": str(d)}, headers=mgr)
    assert r.status_code == 409 and "on duty" in r.json()["detail"]
    r = sec.post(f"{API}/swaps", json={"requester_id": ms[0], "counterpart_id": ms[1], "work_date": "2027-01-05"}, headers=mgr)
    assert r.status_code == 409 and "No published roster" in r.json()["detail"]
    # a swap whose roster changed underneath is re-checked at approval
    a, b = find_pair(sec, mgr, v["id"], d)
    sid = sec.post(f"{API}/swaps", json={"requester_id": a, "counterpart_id": b, "work_date": str(d)}, headers=mgr).json()["swap_id"]
    assert sec.get(f"{API}/swaps", params={"status": "pending_manager"}, headers=mgr).json()[0]["swap_id"] == sid
    sec.post(f"{API}/leaves", json={"employee_id": b, "leave_type": "Casual", "start_date": str(d), "end_date": str(d)}, headers=mgr)
    lid = sec.get(f"{API}/leaves", headers=mgr).json()["items"][0]["application_id"]
    sec.post(f"{API}/leaves/{lid}/approve", headers=mgr)  # b is now on leave -> new draft exists, published roster unchanged
    pub2 = sec.get(f"{API}/roster/versions", params={"status": "draft"}, headers=mgr).json()
    assert pub2, "leave approval must have produced a draft"
    sec.post(f"{API}/roster/versions/{pub2[0]['id']}/approve", json={"acknowledge_shortage": True}, headers=mgr)
    sec.post(f"{API}/roster/versions/{pub2[0]['id']}/publish", json={"acknowledge_shortage": True}, headers=mgr)
    r = sec.post(f"{API}/swaps/{sid}/approve", headers=mgr)
    assert r.status_code == 409 and "on duty" in r.json()["detail"]
    assert sec.post(f"{API}/swaps/{sid}/reject", json={"note": "b is on leave"}, headers=mgr).json()["status"] == "rejected"


# ------------------------------------------------------ audit & export --

def test_audit_history_filters_and_actor(sec):
    mgr, emp = hdr(sec, "mgr1"), hdr(sec, "emp1")
    v = gen(sec, mgr)
    sec.post(f"{API}/roster/versions/{v['id']}/approve", json={"note": "ok"}, headers=mgr)
    sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr)
    log = sec.get(f"{API}/roster/audit-log", params={"version_id": v["id"]}, headers=mgr).json()
    assert [x["action"] for x in log][::-1] == ["roster_generated", "roster_approved", "roster_published"]
    assert {x["actor"] for x in log} == {"mgr1"}
    assert log[0]["details"]["approved_by"] == "mgr1" and log[1]["details"]["note"] == "ok"
    only = sec.get(f"{API}/roster/audit-log", params={"action": "login"}, headers=mgr).json()
    assert {x["actor"] for x in only} >= {"mgr1", "emp1"}
    assert len(sec.get(f"{API}/roster/audit-log", params={"actor": "emp1", "limit": 1}, headers=mgr).json()) == 1
    assert sec.get(f"{API}/roster/audit-log", params={"since": "2999-01-01T00:00:00"}, headers=mgr).json() == []
    assert sec.get(f"{API}/roster/audit-log", params={"offset": 1000}, headers=mgr).json() == []
    # employee changes are audited too, with before/after
    sec.patch(f"{API}/employees/T003", json={"home_thana": "Mirpur"}, headers=mgr)
    ch = sec.get(f"{API}/roster/audit-log", params={"action": "employee_updated"}, headers=mgr).json()[0]
    assert ch["entity_id"] == "T003" and ch["details"]["after"] == {"home_thana": "Mirpur"}


def test_excel_export_content_and_protection(sec):
    mgr, emp = hdr(sec, "mgr1"), hdr(sec, "emp1")
    v = gen(sec, mgr)
    x = sec.get(f"{API}/roster/versions/{v['id']}/export.xlsx", headers=mgr)
    assert x.status_code == 200 and "_DRAFT.xlsx" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Info", "Roster", "Coverage", "Fairness", "Audit"]
    info = {r[0].value: r[1].value for r in wb["Info"].iter_rows()}
    assert info["Status"].startswith("NOT OFFICIAL") and info["Exported by"] == "mgr1"
    ws = wb["Roster"]
    assert ws.max_row == 71 and ws.max_column == 3 + 14 and ws["A2"].value == "T001"
    assert wb["Audit"].max_row >= 2

    sec.post(f"{API}/roster/versions/{v['id']}/approve", json={"note": "ok"}, headers=mgr)
    sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr)
    x = sec.get(f"{API}/roster/versions/{v['id']}/export.xlsx", headers=mgr)
    assert "DRAFT" not in x.headers["content-disposition"]
    info = {r[0].value: r[1].value for r in load_workbook(io.BytesIO(x.content))["Info"].iter_rows()}
    assert info["Status"] == "OFFICIAL - published" and info["Approved by"].startswith("mgr1")
    assert sec.get(f"{API}/roster/versions/{v['id']}/export.xlsx", headers=emp).status_code == 403
    assert sec.get(f"{API}/roster/audit-log", params={"action": "roster_exported"}, headers=mgr).json()


def test_leave_flow_end_to_end_under_auth(sec):
    """Employee applies, manager previews+approves, draft goes through review and publish."""
    mgr, emp = hdr(sec, "mgr1"), hdr(sec, "emp1")
    v = published(sec, mgr)
    i = (date(2026, 11, 4) - START).days
    who = next(e for e, c in grid(sec, mgr, v["id"]).items() if c[i] in ("M", "E", "N") and e == "T001") \
        if grid(sec, mgr, v["id"])["T001"][i] in ("M", "E", "N") else None
    day = "2026-11-04" if who else next(str(START + timedelta(days=k)) for k in range(7, 14)
                                          if grid(sec, mgr, v["id"])["T001"][k] in ("M", "E", "N"))
    lid = sec.post(f"{API}/leaves", json={"employee_id": "T001", "leave_type": "Sick", "start_date": day, "end_date": day},
                   headers=emp).json()["application_id"]
    pv = sec.post(f"{API}/leaves/{lid}/preview", params={"time_limit_seconds": 6}, headers=mgr).json()
    assert pv["simulated"] and pv["impact"]["removed_assignments"]
    res = sec.post(f"{API}/leaves/{lid}/approve", headers=mgr).json()
    draft = res["roster_version"]
    assert draft["status"] == "draft" and draft["created_by"] == "mgr1"
    assert sec.post(f"{API}/roster/versions/{draft['id']}/approve", headers=mgr).status_code == 200
    assert sec.post(f"{API}/roster/versions/{draft['id']}/publish", headers=mgr).status_code == 200
    mine = grid(sec, emp, draft["id"])["T001"]  # the employee sees the published result
    assert mine[(date.fromisoformat(day) - START).days] == "SL"


# ------------------------------------------------ review workflow details --

def test_edit_needs_reason_and_is_audited_with_old_and_new_value(sec):
    mgr = hdr(sec, "mgr1")
    v = gen(sec, mgr)
    codes = grid(sec, mgr, v["id"])
    i = 8  # Nov 3
    emp = next(e for e, c in codes.items() if c[i] == "M")
    base = {"employee_id": emp, "work_date": "2026-11-03", "code": "R"}
    url = f"{API}/roster/versions/{v['id']}/assignments"
    assert sec.put(url, json=base, headers=mgr).status_code == 422                      # reason missing
    assert sec.put(url, json={**base, "reason": "x"}, headers=mgr).status_code == 422   # too short
    assert sec.put(url, json={**base, "reason": "covering a request"}, headers=mgr).status_code == 200
    log = sec.get(f"{API}/roster/audit-log", params={"action": "roster_cell_edited"}, headers=mgr).json()[0]
    assert log["actor"] == "mgr1" and log["created_at"]
    assert log["details"]["from"] == "M" and log["details"]["to"] == "R" and log["details"]["reason"] == "covering a request"
    assert log["details"]["employee_id"] == emp


def test_manual_edit_cannot_break_eligibility_or_rest_rule(sec):
    mgr = hdr(sec, "mgr1")
    v = gen(sec, mgr)
    codes = grid(sec, mgr, v["id"])
    url = f"{API}/roster/versions/{v['id']}/assignments"
    s = sec.Session()
    from app.models.employee import Employee
    ineligible = next(e.employee_id for e in s.query(Employee) if not e.is_shift_eligible_night)
    s.close()
    day = next(k for k in range(1, 13) if codes[ineligible][k] in ("M", "E"))
    r = sec.put(url, json={"employee_id": ineligible, "work_date": str(START + timedelta(days=day)), "code": "N", "reason": "try"}, headers=mgr)
    assert r.status_code == 409 and "not eligible" in r.json()["detail"]
    s = sec.Session()
    elig = next(e.employee_id for e in s.query(Employee) if e.is_shift_eligible_night)
    s.close()
    k = next((k for k in range(1, 13) if codes[elig][k - 1] == "N" and codes[elig][k] in ("E", "M", "R")), None)
    if k is not None:  # Night yesterday -> Morning today is refused
        r = sec.put(url, json={"employee_id": elig, "work_date": str(START + timedelta(days=k)), "code": "M", "reason": "try"}, headers=mgr)
        assert r.status_code == 409 and "Night followed by Morning" in r.json()["detail"]


def test_review_warnings_missing_employee_and_stale_leave(sec):
    mgr = hdr(sec, "mgr1")
    v = gen(sec, mgr)
    det = sec.get(f"{API}/roster/versions/{v['id']}", headers=mgr).json()
    assert det["warnings"] == []
    # an employee added after generation is missing from the draft
    assert sec.post(f"{API}/employees?auto_roster=false", json={"employee_id": "LATE1", "full_name": "Late Joiner",
                    "weekly_off_day": "Sunday"}, headers=mgr).status_code == 201
    w = sec.get(f"{API}/roster/versions/{v['id']}", headers=mgr).json()["warnings"]
    assert [x["type"] for x in w] == ["missing_employee"] and w[0]["employee_ids"] == ["LATE1"]
    r = sec.post(f"{API}/roster/versions/{v['id']}/approve", headers=mgr)
    assert r.status_code == 409 and "not in this roster" in r.json()["detail"]
    # leave approved without roster sync -> this draft predates it
    emp = next(e for e, c in grid(sec, mgr, v["id"]).items() if c[9] in ("M", "E", "N"))
    lid = sec.post(f"{API}/leaves", json={"employee_id": emp, "leave_type": "Casual", "start_date": "2026-11-04",
                   "end_date": "2026-11-04"}, headers=mgr).json()["application_id"]
    sec.post(f"{API}/leaves/{lid}/approve", json={"auto_roster": False}, headers=mgr)
    types = {x["type"] for x in sec.get(f"{API}/roster/versions/{v['id']}", headers=mgr).json()["warnings"]}
    assert types == {"missing_employee", "leave_not_applied"}
    # re-optimising picks up both; warnings are gone on the new draft
    r = sec.post(f"{API}/roster/regenerate", json={"base_version_id": v["id"], "from_date": "2026-10-26", **Q}, headers=mgr)
    assert r.status_code == 201, r.text
    assert sec.get(f"{API}/roster/versions/{r.json()['id']}", headers=mgr).json()["warnings"] == []
    assert sec.get(f"{API}/roster/versions/{v['id']}", headers=hdr(sec, "emp1")).status_code == 403
    assert "warnings" not in sec.get(f"{API}/roster/versions/{published(sec, mgr)['id']}", headers=hdr(sec, "emp1")).json()


def test_change_after_publish_needs_reason_and_makes_new_version(sec):
    mgr = hdr(sec, "mgr1")
    v = published(sec, mgr)
    body = {"base_version_id": v["id"], "from_date": "2026-11-02", **Q}
    r = sec.post(f"{API}/roster/regenerate", json=body, headers=mgr)
    assert r.status_code == 422 and "reason" in r.json()["detail"]
    r = sec.post(f"{API}/roster/regenerate", json={**body, "reason": "demand change"}, headers=mgr)
    assert r.status_code == 201 and r.json()["status"] == "draft" and r.json()["reason"] == "demand change"
    assert r.json()["base_version_id"] == v["id"] and r.json()["created_by"] == "mgr1"
    # the published version stays official until the new draft is approved and published
    assert sec.get(f"{API}/roster/versions/{v['id']}", headers=hdr(sec, "emp1")).json()["version"]["status"] == "published"


def test_version_history_shows_old_versions_and_who_did_what(sec):
    mgr, mgr2 = hdr(sec, "mgr1"), hdr(sec, "mgr2")
    v1 = published(sec, mgr)
    r = sec.post(f"{API}/roster/regenerate", json={"base_version_id": v1["id"], "from_date": "2026-11-02",
                 "reason": "demand change", **Q}, headers=mgr2).json()
    sec.post(f"{API}/roster/versions/{r['id']}/reject", json={"note": "too many nights"}, headers=mgr)
    h = sec.get(f"{API}/roster/versions/{v1['id']}/history", headers=mgr).json()
    vs = {x["version_no"]: x for x in h["versions"]}
    assert vs[1]["status"] == "published" and vs[1]["created_by"] == "mgr1" and vs[1]["approved_by"] == "mgr1" and vs[1]["published_by"] == "mgr1"
    assert vs[2]["status"] == "rejected" and vs[2]["created_by"] == "mgr2" and vs[2]["review_note"] == "too many nights"
    assert vs[2]["base_version_id"] == v1["id"] and vs[2]["reason"] == "demand change"
    acts = [(e["action"], e["actor"]) for e in h["events"]]
    assert ("roster_rejected", "mgr1") in acts and ("roster_regenerated", "mgr2") in acts and ("roster_published", "mgr1") in acts
    assert all(e["time"] for e in h["events"])
    assert sec.get(f"{API}/roster/versions/{v1['id']}/history", headers=hdr(sec, "emp1")).status_code == 403
    # the old version can still be opened and diffed
    assert sec.get(f"{API}/roster/versions/{v1['id']}/diff", params={"against": r["id"]}, headers=mgr).status_code == 200
