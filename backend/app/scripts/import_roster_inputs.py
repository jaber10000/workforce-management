"""Import Daily_Requirement and Leave_Applications sheets from a roster workbook.

Usage from backend directory:
    python -m app.scripts.import_roster_inputs path/to/workbook.xlsx
(Employees are imported separately with app.scripts.import_employees.)
"""
import argparse
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.core.leave_policy import LEAVE_TYPES
from app.models.leave import LeaveApplication
from app.services.roster_service import upsert_requirements


def _d(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def import_inputs(path, db=None, commit: bool = True) -> dict:
    own = db is None
    if own:
        from app.db.session import SessionLocal, Base, engine
        Base.metadata.create_all(bind=engine)
        db = SessionLocal()
    wb = load_workbook(path, data_only=True, read_only=True)
    result = {"requirements": 0, "leaves_created": 0, "leaves_updated": 0, "errors": []}
    try:
        if "Daily_Requirement" in wb.sheetnames:
            items = []
            for i, r in enumerate(wb["Daily_Requirement"].iter_rows(min_row=2, values_only=True), start=2):
                d = _d(r[0])
                if d is None:
                    continue
                try:
                    items.append({"work_date": d, "morning_required": int(r[2]), "evening_required": int(r[3]),
                                  "night_required": int(r[4]), "min_senior": int(r[5] if r[5] is not None else 3)})
                except (TypeError, ValueError) as exc:
                    result["errors"].append(f"Daily_Requirement row {i}: {exc}")
            result["requirements"] = upsert_requirements(db, items)
        if "Leave_Applications" in wb.sheetnames:
            for i, r in enumerate(wb["Leave_Applications"].iter_rows(min_row=2, values_only=True), start=2):
                if not r[0]:
                    continue
                try:
                    app_id, emp, kind = str(r[0]).strip(), str(r[1]).strip(), str(r[2]).strip().capitalize()
                    if kind not in LEAVE_TYPES:
                        raise ValueError(f"unknown leave type {kind}")
                    half = str(r[5]).strip().lower() in {"yes", "y", "true", "1"}
                    status = str(r[7]).strip().capitalize() if r[7] else "Pending"
                    start, end = _d(r[3]), _d(r[4])
                    if not start or not end:
                        raise ValueError("invalid dates")
                    days = float(r[6]) if r[6] is not None else (0.5 if half else (end - start).days + 1)
                    row = db.scalar(select(LeaveApplication).where(LeaveApplication.application_id == app_id))
                    vals = dict(employee_id=emp, leave_type=kind, start_date=start, end_date=end,
                                half_day=half, days=days, status=status, reason=str(r[8]) if r[8] else None)
                    if row:
                        for k, v in vals.items():
                            setattr(row, k, v)
                        result["leaves_updated"] += 1
                    else:
                        db.add(LeaveApplication(application_id=app_id, **vals))
                        result["leaves_created"] += 1
                except Exception as exc:  # keep importing the rest
                    result["errors"].append(f"Leave_Applications row {i}: {exc}")
        if commit:
            db.commit()
        else:
            db.flush()
    except Exception:
        if commit:
            db.rollback()
        raise
    finally:
        wb.close()
        if own:
            db.close()
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    if not args.workbook.exists():
        raise SystemExit(f"File not found: {args.workbook}")
    res = import_inputs(args.workbook)
    print(f"Requirements: {res['requirements']}; leaves created: {res['leaves_created']}, "
          f"updated: {res['leaves_updated']}; errors: {len(res['errors'])}")
    for e in res["errors"]:
        print(e)


if __name__ == "__main__":
    main()
