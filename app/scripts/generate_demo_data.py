"""Generate 5 demo roster workbooks for demos and tests (no real Excel needed).

Usage from backend directory:
    python -m app.scripts.generate_demo_data [output_dir]

Default output dir: ../demo_data

Each workbook contains:
    README              - what the scenario tests
    Employee_Master     - importable by app.scripts.import_employees
    Daily_Requirement   - M/E/N manpower targets + minimum senior per shift
    Leave_Applications  - Sick / Casual / Annual leave (full + half day)
    Roster_Draft        - 14-day draft roster (simple greedy demo, NOT the OR-Tools engine)
    Coverage_Report     - per day/shift shortage report
    Fairness_Summary    - per-employee shift counts

All people, phone numbers and IDs are fictional.
"""
from __future__ import annotations

import random
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROSTER_START = date(2026, 10, 26)  # Monday
ROSTER_DAYS = 14
WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

FIRST_NAMES = [
    "Abdul", "Rahim", "Karim", "Hasan", "Hossain", "Jamal", "Kamal", "Nasir", "Rafiq", "Sabbir",
    "Tanvir", "Imran", "Faisal", "Mahmud", "Shakil", "Arif", "Mehedi", "Rubel", "Sohel", "Masud",
    "Fatema", "Ayesha", "Nusrat", "Sabina", "Rokeya", "Shirin", "Taslima", "Nasrin", "Farhana", "Mim",
    "Sumaiya", "Rina", "Lima", "Joya", "Mitu", "Popy", "Shila", "Dolly", "Anika", "Tania",
]
LAST_NAMES = [
    "Ahmed", "Khan", "Chowdhury", "Islam", "Rahman", "Uddin", "Sarker", "Hossain", "Mia", "Akter",
    "Begum", "Khatun", "Talukder", "Biswas", "Mondal", "Sheikh", "Howlader", "Bhuiyan", "Miah", "Sikder",
]
THANAS = [
    "Mirpur", "Uttara", "Mohammadpur", "Dhanmondi", "Gulshan", "Banani", "Badda", "Rampura",
    "Khilgaon", "Motijheel", "Jatrabari", "Demra", "Savar", "Keraniganj", "Tejgaon", "Pallabi",
    "Kafrul", "Lalbagh", "Sutrapur", "Tongi",
]
ROLE_LABEL = {"general": "General", "senior": "Senior Support", "supervisor": "Supervisor"}


@dataclass
class Scenario:
    key: str
    title: str
    purpose: str
    employees: int = 200
    seniors: int = 24
    supervisors: int = 8
    night_eligible: int = 90
    need_m: int = 55
    need_e: int = 45
    need_n: int = 14
    min_senior: int = 3
    leaves: int = 6
    leave_cluster: bool = False
    new_joiners: int = 0
    inactive: int = 0
    seed: int = 1


SCENARIOS = [
    Scenario(
        "01_standard", "Standard balanced team",
        "Baseline happy path: 200 employees, enough night-eligible staff, few leaves. "
        "Expect zero shortages and a tight night-shift fairness spread.",
        seed=101,
    ),
    Scenario(
        "02_night_shortage", "Night shift shortage",
        "Only 15 night-eligible employees against a night target of 14. Weekly offs, rest rules "
        "and leave should create Night shortages in the Coverage_Report. Tests shortage reporting.",
        night_eligible=15, leaves=8, seed=202,
    ),
    Scenario(
        "03_senior_tight", "Tight senior coverage",
        "Only 9 senior-support staff while each shift needs 3 seniors. Senior coverage "
        "shortfalls should appear. Tests the senior-support constraint.",
        seniors=9, supervisors=6, seed=303,
    ),
    Scenario(
        "04_leave_heavy", "Leave heavy fortnight",
        "45 leave applications (sick/casual/annual, some half-day, some pending) clustered in "
        "the second week. Tests leave -> roster sync, replacement and re-optimization.",
        leaves=45, leave_cluster=True, seed=404,
    ),
    Scenario(
        "05_new_joiners", "New joiners and inactive staff",
        "215 records: 15 new joiners (some join mid-roster, none night-eligible yet) and 6 inactive "
        "employees. Tests 'new employee added -> coverage re-check' and inactive filtering.",
        employees=215, new_joiners=15, inactive=6, leaves=8, seed=505,
    ),
]


@dataclass
class Emp:
    emp_id: str
    name: str
    gender: str
    phone: str
    thana: str
    role: str  # general | senior | supervisor
    off_day: str
    half_day: str | None
    joining: date
    morning: bool
    evening: bool
    night: bool
    active: bool
    notes: str = ""
    # roster bookkeeping
    counts: dict = field(default_factory=lambda: defaultdict(int))
    consec_night: int = 0
    prev: str = ""


@dataclass
class Leave:
    emp_id: str
    kind: str
    start: date
    end: date
    half_day: bool
    status: str
    reason: str


def build_employees(sc: Scenario, rng: random.Random) -> list[Emp]:
    total = sc.employees
    roles = (["senior"] * sc.seniors) + (["supervisor"] * sc.supervisors)
    roles += ["general"] * (total - len(roles))
    rng.shuffle(roles)

    new_idx = set(range(total - sc.new_joiners, total))
    inactive_idx = set(rng.sample(range(total - sc.new_joiners), sc.inactive)) if sc.inactive else set()

    # night-eligible only among established, active employees
    candidates = [i for i in range(total) if i not in new_idx and i not in inactive_idx]
    sen_c = [i for i in candidates if roles[i] == "senior"]
    oth_c = [i for i in candidates if roles[i] != "senior"]
    n_sen = min(len(sen_c), max(sc.min_senior * 2, sc.night_eligible // 6))
    night_idx = set(rng.sample(sen_c, n_sen))
    night_idx |= set(rng.sample(oth_c, min(sc.night_eligible - n_sen, len(oth_c))))

    used = set()
    emps: list[Emp] = []
    for i in range(total):
        while True:
            name = f"{rng.choice(FIRST_NAMES)} {rng.choice(LAST_NAMES)}"
            if name not in used:
                used.add(name)
                break
        female_names = {"Fatema", "Ayesha", "Nusrat", "Sabina", "Rokeya", "Shirin", "Taslima", "Nasrin",
                        "Farhana", "Mim", "Sumaiya", "Rina", "Lima", "Joya", "Mitu", "Popy", "Shila",
                        "Dolly", "Anika", "Tania"}
        gender = "Female" if name.split()[0] in female_names else "Male"
        off = WEEKDAYS[i % 7]
        half = rng.choice(WEEKDAYS) if rng.random() < 0.12 else None
        if half == off:
            half = None
        if i in new_idx:
            # first half joined before roster start, the rest join during the roster window
            joining = ROSTER_START - timedelta(days=rng.randint(3, 20)) if (i % 2 == 0) \
                else ROSTER_START + timedelta(days=rng.randint(2, 9))
            notes = "New joiner - onboarding"
        else:
            joining = date(rng.randint(2016, 2025), rng.randint(1, 12), rng.randint(1, 28))
            notes = ""
        is_night = i in night_idx
        emps.append(Emp(
            emp_id=f"EMP{i + 1:03d}", name=name, gender=gender,
            phone=f"01700-{100000 + i:06d}", thana=rng.choice(THANAS), role=roles[i],
            off_day=off, half_day=half, joining=joining,
            morning=True, evening=rng.random() > 0.05, night=is_night,
            active=i not in inactive_idx,
            notes=("Inactive - resigned" if i in inactive_idx else notes),
        ))
    return emps


def build_leaves(sc: Scenario, emps: list[Emp], rng: random.Random) -> list[Leave]:
    pool = [e for e in emps if e.active and e.joining <= ROSTER_START]
    leaves: list[Leave] = []
    reasons = {"Sick": "Fever / medical", "Casual": "Personal work", "Annual": "Family trip"}
    taken: dict[str, list[tuple[date, date]]] = defaultdict(list)
    while len(leaves) < sc.leaves:
        e = rng.choice(pool)
        kind = rng.choices(["Sick", "Casual", "Annual"], weights=[3, 4, 2])[0]
        if sc.leave_cluster:
            start = ROSTER_START + timedelta(days=rng.randint(7, 11))
        else:
            start = ROSTER_START + timedelta(days=rng.randint(0, ROSTER_DAYS - 3))
        length = {"Sick": rng.randint(1, 3), "Casual": rng.randint(1, 2), "Annual": rng.randint(3, 6)}[kind]
        half = kind in ("Casual", "Sick") and length == 1 and rng.random() < 0.3
        end = start + timedelta(days=length - 1)
        if any(start <= b and a <= end for a, b in taken[e.emp_id]):
            continue  # an employee never has overlapping leave applications
        taken[e.emp_id].append((start, end))
        status = "Approved" if rng.random() < 0.8 else "Pending"
        leaves.append(Leave(e.emp_id, kind, start, end, half, status, reasons[kind]))
    return leaves


# ---------------------------------------------------------------- roster ---

def generate_roster(sc: Scenario, emps: list[Emp], leaves: list[Leave], rng: random.Random):
    """Simple greedy roster used only to give demo data a realistic draft.

    Rules: weekly off, approved leave, joining date, shift eligibility, no Night->Morning
    next day, max 3 consecutive nights, half-duty/half-day = 0.5
    weight (never Night), seniors first, least-loaded employee first (fairness).
    """
    days = [ROSTER_START + timedelta(days=i) for i in range(ROSTER_DAYS)]
    leave_map: dict[tuple[str, date], Leave] = {}
    for lv in leaves:
        if lv.status != "Approved":
            continue
        d = lv.start
        while d <= lv.end:
            leave_map[(lv.emp_id, d)] = lv
            d += timedelta(days=1)

    grid: dict[str, dict[date, str]] = {e.emp_id: {} for e in emps}
    coverage = []
    leave_code = {"Sick": "SL", "Casual": "CL", "Annual": "AL"}

    for day in days:
        wd = WEEKDAYS[day.weekday()]
        avail: list[Emp] = []
        weight: dict[str, float] = {}
        for e in emps:
            if not e.active:
                grid[e.emp_id][day] = "-"
                e.prev = ""
                continue
            if day < e.joining:
                grid[e.emp_id][day] = "-"
                e.prev = ""
                continue
            if wd == e.off_day:
                grid[e.emp_id][day] = "WO"
                e.prev = "WO"
                e.consec_night = 0
                continue
            lv = leave_map.get((e.emp_id, day))
            if lv and not lv.half_day:
                grid[e.emp_id][day] = leave_code[lv.kind]
                e.prev = "L"
                e.consec_night = 0
                continue
            w = 0.5 if (wd == e.half_day or (lv and lv.half_day)) else 1.0
            weight[e.emp_id] = w
            avail.append(e)

        assigned: dict[str, str] = {}
        got = {"M": 0.0, "E": 0.0, "N": 0.0}
        senior_got = {"M": 0, "E": 0, "N": 0}
        need = {"M": sc.need_m, "E": sc.need_e, "N": sc.need_n}

        def allowed(e: Emp, shift: str) -> bool:
            if e.emp_id in assigned:
                return False
            if shift == "N":
                return e.night and weight[e.emp_id] == 1.0 and e.consec_night < 3
            if shift == "E":
                return e.evening
            return e.morning and e.prev != "N"

        def load_key(e: Emp, shift: str):
            return (e.counts[shift], sum(e.counts[s] for s in "MEN"), rng.random())

        def take(e: Emp, shift: str):
            assigned[e.emp_id] = shift
            got[shift] += weight[e.emp_id]
            if e.role == "senior":
                senior_got[shift] += 1

        for shift in ("N", "E", "M"):  # scarcest shift first
            seniors = sorted([e for e in avail if e.role == "senior" and allowed(e, shift)],
                             key=lambda x: load_key(x, shift))
            for e in seniors:
                if senior_got[shift] >= sc.min_senior:
                    break
                take(e, shift)
            rest = sorted([e for e in avail if allowed(e, shift)], key=lambda x: load_key(x, shift))
            for e in rest:
                if got[shift] >= need[shift]:
                    break
                take(e, shift)

        for e in avail:
            shift = assigned.get(e.emp_id)
            if shift:
                half = weight[e.emp_id] == 0.5
                grid[e.emp_id][day] = f"{shift}-H" if half else shift
                e.counts[shift] += 1
                e.consec_night = e.consec_night + 1 if shift == "N" else 0
                e.prev = shift
            else:
                grid[e.emp_id][day] = "R"  # reserve / standby
                e.prev = "R"
                e.consec_night = 0

        coverage.append((day, wd, got, need, senior_got))
    return days, grid, coverage


# -------------------------------------------------------------- workbook ---

HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")
FILLS = {
    "M": "FFF2CC", "E": "FCE4D6", "N": "D9D2E9", "WO": "D9D9D9", "R": "E2EFDA",
    "SL": "F8CBAD", "CL": "BDD7EE", "AL": "C6E0B4", "-": "F2F2F2",
}
BAD_FILL = PatternFill("solid", fgColor="FF9999")


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autosize(ws, min_w=8, max_w=40):
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_w, min(max_w, width))


def yes_no(v: bool) -> str:
    return "Yes" if v else "No"


def write_workbook(sc: Scenario, out_dir: Path) -> Path:
    rng = random.Random(sc.seed)
    emps = build_employees(sc, rng)
    leaves = build_leaves(sc, emps, rng)
    days, grid, coverage = generate_roster(sc, emps, leaves, rng)

    wb = Workbook()

    # README
    ws = wb.active
    ws.title = "README"
    ws.append(["Demo scenario", sc.title])
    ws.append(["Key", sc.key])
    ws.append(["Purpose", sc.purpose])
    ws.append(["Roster window", f"{days[0].isoformat()} to {days[-1].isoformat()} ({ROSTER_DAYS} days)"])
    ws.append(["Employees", len(emps)])
    ws.append(["Targets (M/E/N)", f"{sc.need_m}/{sc.need_e}/{sc.need_n}; min senior per shift: {sc.min_senior}"])
    ws.append([])
    ws.append(["Shift codes", "M=Morning, E=Evening, N=Night, X-H=half duty, R=reserve, WO=weekly off, "
                              "SL/CL/AL=sick/casual/annual leave, '-'=not employed / inactive"])
    ws.append(["Note", "All data is FICTIONAL. Roster_Draft is a simple greedy demo, not the OR-Tools engine."])
    for r in ws.iter_rows(min_row=1, max_col=1):
        r[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    # Employee_Master
    ws = wb.create_sheet("Employee_Master")
    ws.append(["Employee ID", "Employee Name", "Gender", "Phone Number (Demo)", "Home Thana", "Role / Pool",
               "Weekly Off Day", "Half-duty Day", "Joining Date", "Morning Eligible", "Evening Eligible",
               "Night Eligible", "Active", "Notes"])
    for e in emps:
        ws.append([e.emp_id, e.name, e.gender, e.phone, e.thana, ROLE_LABEL[e.role], e.off_day,
                   e.half_day or "", e.joining, yes_no(e.morning), yes_no(e.evening), yes_no(e.night),
                   yes_no(e.active), e.notes])
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        row[0].number_format = "yyyy-mm-dd"
    style_header(ws)
    ws.freeze_panes = "C2"
    autosize(ws)

    # Daily_Requirement
    ws = wb.create_sheet("Daily_Requirement")
    ws.append(["Date", "Weekday", "Morning Required", "Evening Required", "Night Required",
               "Min Senior per Shift"])
    for d in days:
        ws.append([d, WEEKDAYS[d.weekday()], sc.need_m, sc.need_e, sc.need_n, sc.min_senior])
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "yyyy-mm-dd"
    style_header(ws)
    autosize(ws)

    # Leave_Applications
    ws = wb.create_sheet("Leave_Applications")
    ws.append(["Application ID", "Employee ID", "Leave Type", "Start Date", "End Date", "Half Day",
               "Days", "Status", "Reason"])
    for i, lv in enumerate(leaves, start=1):
        days_n = 0.5 if lv.half_day else (lv.end - lv.start).days + 1
        ws.append([f"LV{i:04d}", lv.emp_id, lv.kind, lv.start, lv.end, yes_no(lv.half_day), days_n,
                   lv.status, lv.reason])
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    style_header(ws)
    autosize(ws)

    # Roster_Draft
    ws = wb.create_sheet("Roster_Draft")
    ws.append(["Employee ID", "Employee Name", "Role"] + [d.strftime("%a %d-%b") for d in days])
    for e in emps:
        ws.append([e.emp_id, e.name, ROLE_LABEL[e.role]] + [grid[e.emp_id][d] for d in days])
    style_header(ws)
    for row in ws.iter_rows(min_row=2, min_col=4):
        for c in row:
            base = str(c.value).split("-")[0] if c.value not in ("-", None) else "-"
            key = c.value if c.value in FILLS else base
            if key in FILLS:
                c.fill = PatternFill("solid", fgColor=FILLS[key])
            c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "D2"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 16
    for i in range(4, 4 + len(days)):
        ws.column_dimensions[get_column_letter(i)].width = 11

    # Coverage_Report
    ws = wb.create_sheet("Coverage_Report")
    ws.append(["Date", "Weekday",
               "Morning Assigned", "Morning Required", "Morning Shortage",
               "Evening Assigned", "Evening Required", "Evening Shortage",
               "Night Assigned", "Night Required", "Night Shortage",
               "Senior M", "Senior E", "Senior N", "Senior Shortage", "Status"])
    for day, wd, got, need, sen in coverage:
        sm = max(0, need["M"] - got["M"])
        se = max(0, need["E"] - got["E"])
        sn = max(0, need["N"] - got["N"])
        sen_short = sum(max(0, sc.min_senior - sen[s]) for s in "MEN")
        status = "OK" if (sm + se + sn + sen_short) == 0 else "SHORTAGE"
        ws.append([day, wd, got["M"], need["M"], sm, got["E"], need["E"], se, got["N"], need["N"], sn,
                   sen["M"], sen["E"], sen["N"], sen_short, status])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        if row[-1].value == "SHORTAGE":
            row[-1].fill = BAD_FILL
            row[-1].font = Font(bold=True)
    style_header(ws)
    autosize(ws)

    # Fairness_Summary
    ws = wb.create_sheet("Fairness_Summary")
    ws.append(["Employee ID", "Employee Name", "Night Eligible", "Morning", "Evening", "Night", "Total Shifts"])
    for e in emps:
        if not e.active:
            continue
        ws.append([e.emp_id, e.name, yes_no(e.night), e.counts["M"], e.counts["E"], e.counts["N"],
                   e.counts["M"] + e.counts["E"] + e.counts["N"]])
    n_counts = [e.counts["N"] for e in emps if e.active and e.night]
    ws.append([])
    ws.append(["", "Night spread (min / max / avg) among night-eligible",
               "", min(n_counts) if n_counts else 0, max(n_counts) if n_counts else 0,
               round(sum(n_counts) / len(n_counts), 2) if n_counts else 0])
    style_header(ws)
    autosize(ws)

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"demo_roster_{sc.key}.xlsx"
    wb.save(path)
    return path


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[3] / "demo_data"
    for sc in SCENARIOS:
        print("wrote", write_workbook(sc, out))


if __name__ == "__main__":
    main()
