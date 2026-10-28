import io
from datetime import date, datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.deps import Principal, actor_name, get_db, get_principal, get_today, manager_required, review_required
from app.models.employee import Employee
from app.models.roster import AuditLog, DailyRequirement, RosterVersion
from app.schemas.roster import (ApproveIn, AuditRead, CellEdit, GenerateIn, PublishIn, RegenerateIn, RejectIn,
                                RequirementIn, RequirementRead, VersionRead)
from app.services import roster_service as svc
from app.services.roster_service import RosterError

router = APIRouter(prefix="/roster", tags=["Roster"], dependencies=[Depends(get_principal)])


def _wrap(fn, *args, **kwargs):
    try:
        return fn(*args, **kwargs)
    except RosterError as exc:
        raise HTTPException(exc.status_code, exc.message)


# ------------------------------------------------------------ requirements --

@router.get("/requirements", response_model=list[RequirementRead])
def get_requirements(start_date: date, end_date: date, db: Session = Depends(get_db)):
    if end_date < start_date or (end_date - start_date).days >= svc.MAX_DAYS:
        raise HTTPException(422, "Invalid date range")
    reqs = svc.load_requirements(db, start_date, end_date)
    return [RequirementRead(work_date=d, morning_required=r.morning, evening_required=r.evening,
                            night_required=r.night, min_senior=r.min_senior) for d, r in reqs.items()]


@router.put("/requirements")
def put_requirements(items: list[RequirementIn], db: Session = Depends(get_db), p: Principal = Depends(manager_required)):
    n = svc.upsert_requirements(db, [i.model_dump() for i in items])
    db.commit()
    return {"updated": n}


# ---------------------------------------------------------------- versions --

@router.post("/generate", response_model=VersionRead, status_code=status.HTTP_201_CREATED)
def generate(body: GenerateIn, db: Session = Depends(get_db), p: Principal = Depends(manager_required)):
    return _wrap(svc.generate_roster, db, start_date=body.start_date, end_date=body.end_date, actor=p.username,
                 requirements=[r.model_dump() for r in body.requirements] if body.requirements else None,
                 time_limit=body.time_limit_seconds, reason=body.reason)


@router.post("/regenerate", response_model=VersionRead, status_code=status.HTTP_201_CREATED)
def regenerate(body: RegenerateIn, db: Session = Depends(get_db), today: date = Depends(get_today),
               p: Principal = Depends(manager_required)):
    return _wrap(svc.regenerate_roster, db, base_version_id=body.base_version_id, from_date=body.from_date, actor=p.username,
                 reason=body.reason, time_limit=body.time_limit_seconds, today=today)


@router.get("/versions", response_model=list[VersionRead])
def list_versions(roster_key: str | None = None, status_: str | None = Query(default=None, alias="status"),
                  limit: int = Query(default=50, ge=1, le=200), db: Session = Depends(get_db),
                  p: Principal = Depends(get_principal)):
    stmt = select(RosterVersion)
    if not p.is_reviewer:  # employees only ever see published rosters
        stmt = stmt.where(RosterVersion.status == "published")
    if status_:
        stmt = stmt.where(RosterVersion.status == status_)
    if roster_key:
        stmt = stmt.where(RosterVersion.roster_key == roster_key)
    return db.scalars(stmt.order_by(RosterVersion.id.desc()).limit(limit)).all()


@router.get("/versions/{version_id}")
def version_detail(version_id: int, include_grid: bool = True, db: Session = Depends(get_db),
                   p: Principal = Depends(get_principal)):
    v = _wrap(svc.get_version, db, version_id)
    if not p.is_reviewer and v.status != "published":
        raise HTTPException(403, "Only published rosters are visible to employees")
    out = {"version": VersionRead.model_validate(v).model_dump(mode="json"),
           "coverage": svc.coverage_report(db, version_id)}
    if p.is_reviewer:
        out["warnings"] = svc.version_warnings(db, version_id)
    if include_grid:
        out["grid"] = svc.roster_grid(db, version_id)
    return out


@router.get("/versions/{version_id}/history")
def history(version_id: int, db: Session = Depends(get_db), p: Principal = Depends(review_required)):
    """All versions of this roster window with who created/approved/rejected/published them, plus the audit trail."""
    return _wrap(svc.version_history, db, version_id)


@router.get("/versions/{version_id}/replacements")
def replacements(version_id: int, work_date: date, shift: str = Query(pattern="^[MEN]$"),
                 limit: int = Query(default=10, ge=1, le=50), db: Session = Depends(get_db),
                 p: Principal = Depends(review_required)):
    """Replacement candidates for a short shift (reserve first, then spare cover, then weekly-off call-in)."""
    return {"version_id": version_id, "date": work_date.isoformat(), "shift": shift,
            "candidates": _wrap(svc.replacement_candidates, db, version_id, work_date, shift, limit)}


@router.get("/versions/{version_id}/coverage")
def coverage(version_id: int, db: Session = Depends(get_db), p: Principal = Depends(review_required)):
    _wrap(svc.get_version, db, version_id)
    return svc.coverage_report(db, version_id)


@router.get("/versions/{version_id}/fairness")
def fairness(version_id: int, db: Session = Depends(get_db), p: Principal = Depends(review_required)):
    _wrap(svc.get_version, db, version_id)
    return svc.fairness_report(db, version_id)


@router.get("/versions/{version_id}/diff")
def diff(version_id: int, against: int | None = None, db: Session = Depends(get_db),
         p: Principal = Depends(review_required)):
    return _wrap(svc.version_diff, db, version_id, against)


@router.post("/versions/{version_id}/approve", response_model=VersionRead)
def approve(version_id: int, body: ApproveIn | None = None, db: Session = Depends(get_db),
            p: Principal = Depends(manager_required)):
    """Manager review: draft -> approved (locked, ready to publish)."""
    body = body or ApproveIn()
    return _wrap(svc.approve, db, version_id, acknowledge_shortage=body.acknowledge_shortage, note=body.note,
                 actor=actor_name(p, body.decided_by))


@router.post("/versions/{version_id}/reject", response_model=VersionRead)
def reject(version_id: int, body: RejectIn, db: Session = Depends(get_db), p: Principal = Depends(manager_required)):
    return _wrap(svc.reject, db, version_id, note=body.note, actor=actor_name(p, body.decided_by))


@router.post("/versions/{version_id}/publish", response_model=VersionRead)
def publish(version_id: int, body: PublishIn | None = None, db: Session = Depends(get_db),
            p: Principal = Depends(manager_required)):
    body = body or PublishIn()
    return _wrap(svc.publish, db, version_id, acknowledge_shortage=body.acknowledge_shortage,
                 actor=actor_name(p, body.decided_by))


def _check_team_scope(db: Session, p: Principal, employee_id: str) -> None:
    """A supervisor linked to an employee profile that has a team may only change that team's shifts."""
    if p.role != "supervisor" or not p.employee_id:
        return
    mine = db.scalar(select(Employee.team).where(Employee.employee_id == p.employee_id))
    if not mine:
        return
    theirs = db.scalar(select(Employee.team).where(Employee.employee_id == employee_id))
    if theirs != mine:
        raise HTTPException(403, f"{employee_id} is not in your team ({mine}); only that team's supervisor or a manager can change this shift")


@router.put("/versions/{version_id}/assignments")
def edit_cell(version_id: int, body: CellEdit, db: Session = Depends(get_db), p: Principal = Depends(review_required)):
    _check_team_scope(db, p, body.employee_id)
    row = _wrap(svc.edit_cell, db, version_id, body.employee_id, body.work_date, body.code,
                 actor=p.username, reason=body.reason)
    return {"employee_id": row.employee_id, "work_date": row.work_date.isoformat(), "code": row.code}


@router.delete("/versions/{version_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_draft(version_id: int, db: Session = Depends(get_db), p: Principal = Depends(manager_required)):
    _wrap(svc.delete_draft, db, version_id, actor=p.username)


@router.get("/versions/{version_id}/export.xlsx")
def export_xlsx(version_id: int, db: Session = Depends(get_db), p: Principal = Depends(manager_required)):
    v = _wrap(svc.get_version, db, version_id)
    grid, cov, fair = svc.roster_grid(db, version_id), svc.coverage_report(db, version_id), svc.fairness_report(db, version_id)
    wb = Workbook()
    info = wb.active
    info.title = "Info"
    state = "OFFICIAL - published" if v.status == "published" else f"NOT OFFICIAL - {v.status.upper()}"
    for row in [
        ("Status", state), ("Roster", f"{v.start_date} to {v.end_date}"), ("Version", f"v{v.version_no} (id {v.id})"),
        ("Trigger", v.trigger), ("Reason", v.reason or ""), ("Created by", v.created_by or ""),
        ("Approved by", f"{v.approved_by or ''} {v.approved_at:%Y-%m-%d %H:%M} UTC" if v.approved_at else ""),
        ("Review note", v.review_note or ""),
        ("Published by", f"{v.published_by or ''} {v.published_at:%Y-%m-%d %H:%M} UTC" if v.published_at else ""),
        ("Shortage (person-shifts)", cov["totals"]["total"]), ("Days with shortage", cov["days_with_shortage"]),
        ("Exported by", p.username), ("Exported at", f"{datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC"),
    ]:
        info.append(list(row))
    info["B1"].font = Font(bold=True, color="B3261E" if v.status != "published" else "1B7F5C")
    ws = wb.create_sheet("Roster")
    ws.append(["Employee ID", "Employee Name", "Role"] + [f"{d['weekday'][:3]} {d['date']}" for d in grid["days"]])
    fills = {"M": "FFF2CC", "E": "FCE4D6", "N": "D9D2E9", "WO": "D9D9D9", "R": "E2EFDA",
             "SL": "F8CBAD", "CL": "F8CBAD", "AL": "F8CBAD"}
    for r in grid["rows"]:
        ws.append([r["employee_id"], r["employee_name"], r["role"]] + r["codes"])
        for cell in ws[ws.max_row][3:]:
            val = str(cell.value)
            fill = fills.get(val[:1] if val[:1] in "MEN" and (len(val) == 1 or val.endswith("-H")) else val)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "D2"
    cs = wb.create_sheet("Coverage")
    cs.append(["Date", "Weekday", "M assigned", "M required", "M short", "E assigned", "E required", "E short",
               "N assigned", "N required", "N short", "Senior short", "Status"])
    for d in cov["days"]:
        m, e, n = d["shifts"]["morning"], d["shifts"]["evening"], d["shifts"]["night"]
        cs.append([d["date"], d["weekday"], m["assigned"], m["required"], m["shortage"], e["assigned"],
                   e["required"], e["shortage"], n["assigned"], n["required"], n["shortage"],
                   d["senior_shortage_total"], d["status"]])
    fs = wb.create_sheet("Fairness")
    fs.append(["Employee ID", "Name", "Role", "Night eligible", "Morning", "Evening", "Night", "Total", "Half duty",
               "Leave days", "Reserve days", "Undesirable score"])
    for r in fair["employees"]:
        fs.append([r["employee_id"], r["employee_name"], r["role"], "Yes" if r["night_eligible"] else "No",
                   r["morning"], r["evening"], r["night"], r["total"], r["half_duty"], r["leave_days"],
                   r["reserve_days"], r["undesirable_score"]])
    au = wb.create_sheet("Audit")
    au.append(["Time (UTC)", "Actor", "Action", "Entity", "Details"])
    for a_ in db.scalars(select(AuditLog).where(AuditLog.version_id == version_id).order_by(AuditLog.id)):
        au.append([f"{a_.created_at:%Y-%m-%d %H:%M:%S}", a_.actor, a_.action, f"{a_.entity_type} {a_.entity_id}",
                   str(a_.details)[:500]])
    for sheet in wb:
        for c in sheet[1]:
            c.font = Font(bold=True)
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 14 if sheet.title == "Roster" and i > 3 else 22
    info.column_dimensions["B"].width = 48
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    suffix = "" if v.status == "published" else f"_{v.status.upper()}"
    name = f"roster_{v.roster_key}_v{v.version_no}{suffix}.xlsx"
    svc.audit.log(db, "roster_exported", "roster_version", v.id, {"status": v.status}, version_id=v.id, actor=p.username)
    db.commit()
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ------------------------------------------------------------------- audit --

@router.get("/audit-log", response_model=list[AuditRead])
def audit_log(version_id: int | None = None, action: str | None = None, actor: str | None = None,
              entity_type: str | None = None, entity_id: str | None = None,
              since: datetime | None = None, until: datetime | None = None,
              offset: int = Query(default=0, ge=0), limit: int = Query(default=100, ge=1, le=500),
              db: Session = Depends(get_db), p: Principal = Depends(manager_required)):
    """Audit history (newest first). Every roster change, approval, leave, swap and login is recorded."""
    stmt = select(AuditLog)
    for col, val in ((AuditLog.version_id, version_id), (AuditLog.action, action), (AuditLog.actor, actor),
                     (AuditLog.entity_type, entity_type), (AuditLog.entity_id, entity_id)):
        if val is not None:
            stmt = stmt.where(col == val)
    if since:
        stmt = stmt.where(AuditLog.created_at >= since)
    if until:
        stmt = stmt.where(AuditLog.created_at <= until)
    return db.scalars(stmt.order_by(AuditLog.id.desc()).offset(offset).limit(limit)).all()
