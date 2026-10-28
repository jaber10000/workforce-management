"""In-app Excel import: employees, daily requirements, leave applications and an existing weekly roster.

One workbook can carry any mix of these sheets (the same layout as the demo workbooks and the app's own exports):

  Employee_Master | Employees     -> employees (upsert by Employee ID)
  Daily_Requirement                -> per-day headcount targets
  Leave_Applications               -> leave applications (upsert by Application ID)
  Weekly_Roster [+ Senior_Roster]  -> an existing roster, stored as a NEW DRAFT version (never touches a published one)

A dry run executes everything inside the transaction and rolls it back, so the preview numbers are exactly what the
real import would do.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.employee import Employee
from app.models.roster import RosterAssignment, RosterVersion
from app.scripts.import_employees import SHEET_NAMES, import_workbook
from app.scripts.import_roster_inputs import import_inputs
from app.services import audit
from app.services.roster_engine import code_weight
from app.services.roster_service import MAX_DAYS, WEEKDAYS, RosterError, coverage_report

MAX_BYTES = 5 * 1024 * 1024
GRID_SHEETS = ("Weekly_Roster", "Senior_Roster", "Roster")
MAX_REPORTED = 50

# Codes seen in real workbooks -> the engine's codes.
CODE_ALIASES = {
    "M": "M", "E": "E", "N": "N", "M-H": "M-H", "E-H": "E-H", "N-H": "N-H",
    "R": "R", "WO": "WO", "OFF": "WO", "SL": "SL", "CL": "CL", "AL": "AL", "-": "-", "": "-",
    "HDLM": "M-H",  # 11:00-15:00
    "HDE": "E-H",   # 15:00-19:00
    "HDLE": "E-H",  # 19:00-23:00
}
_DATE_IN_HEADER = re.compile(r"(\d{4})-(\d{2})-(\d{2})")


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        m = _DATE_IN_HEADER.search(value)
        if m:
            try:
                return date(int(m[1]), int(m[2]), int(m[3]))
            except ValueError:
                return None
    return None


def _weekday_index(value) -> int | None:
    if not isinstance(value, str):
        return None
    t = value.strip().lower()
    for i, name in enumerate(WEEKDAYS):
        if t == name.lower() or t == name[:3].lower():
            return i
    return None


def _grid_columns(header: tuple, week_start: date | None):
    """Map each header cell to a date. Date headers win; weekday names need week_start."""
    cols: dict[int, date] = {}
    for i, h in enumerate(header):
        d = _as_date(h)
        if d is not None:
            cols[i] = d
    if cols:
        return cols
    if week_start is None:
        return {}
    for i, h in enumerate(header):
        wd = _weekday_index(h)
        if wd is not None:
            cols[i] = week_start + timedelta(days=(wd - week_start.weekday()) % 7)
    return cols


def parse_grid(wb, week_start: date | None, errors: list[str]) -> dict[tuple[str, date], str]:
    """Read every grid sheet into {(employee_id, date): code}. Later sheets do not override earlier ones."""
    cells: dict[tuple[str, date], str] = {}
    found = False
    for name in GRID_SHEETS:
        if name not in wb.sheetnames:
            continue
        rows = wb[name].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        found = True
        id_col = next((i for i, h in enumerate(header) if str(h or "").strip().lower() == "employee id"), None)
        cols = _grid_columns(header, week_start)
        if id_col is None:
            errors.append(f"{name}: no 'Employee ID' column")
            continue
        if not cols:
            errors.append(f"{name}: no day columns found" + ("" if week_start else " (choose the roster start date)"))
            continue
        for n, row in enumerate(rows, start=2):
            emp = str(row[id_col]).strip() if id_col < len(row) and row[id_col] is not None else ""
            if not emp:
                continue
            for i, d in cols.items():
                raw = row[i] if i < len(row) else None
                code = CODE_ALIASES.get(str(raw).strip().upper() if raw is not None else "")
                if code is None:
                    errors.append(f"{name} row {n} {emp} {d}: unknown shift code '{raw}' (left empty)")
                    code = "-"
                cells.setdefault((emp, d), code)
    if not found:
        return {}
    return cells


def _create_draft(db: Session, cells: dict[tuple[str, date], str], *, actor: str, reason: str | None,
                  filename: str | None, warnings: list[str]) -> tuple[RosterVersion, dict]:
    known = {e.employee_id: e for e in db.scalars(select(Employee))}
    unknown = sorted({emp for emp, _ in cells if emp not in known})
    cells = {k: c for k, c in cells.items() if k[0] in known}
    if not cells:
        raise RosterError("None of the employee IDs in the roster sheet exist - import employees first", 422)
    start, end = min(d for _, d in cells), max(d for _, d in cells)
    if (end - start).days + 1 > MAX_DAYS:
        raise RosterError(f"Roster window is limited to {MAX_DAYS} days", 422)
    night_ineligible = sorted({emp for (emp, _), code in cells.items()
                               if code.startswith("N") and not known[emp].is_shift_eligible_night})
    if night_ineligible:
        warnings.append(f"{len(night_ineligible)} employee(s) have Night shifts but are not marked night-eligible "
                        f"in their record (e.g. {', '.join(night_ineligible[:5])}). Add a 'Night Eligible' column "
                        "to Employee_Master or fix them in Employees.")
    key = f"{start.isoformat()}_{end.isoformat()}"
    version_no = (db.scalar(select(func.max(RosterVersion.version_no)).where(RosterVersion.roster_key == key)) or 0) + 1
    v = RosterVersion(roster_key=key, version_no=version_no, start_date=start, end_date=end, status="draft",
                      trigger="import", reason=reason or "Imported from Excel", created_by=actor,
                      solver_status="imported", summary={})
    db.add(v)
    db.flush()
    db.bulk_insert_mappings(RosterAssignment, [
        {"version_id": v.id, "employee_id": emp, "work_date": d, "code": code, "weight": code_weight(code)}
        for (emp, d), code in cells.items()])
    db.flush()
    active = [e for e in known.values() if e.is_active]
    cov = coverage_report(db, v.id, employees=active)
    missing = sorted(e.employee_id for e in active if not any((e.employee_id, d) in cells for d in (start, end)))
    v.summary = {"employees": len({e for e, _ in cells}), "shortage": cov["totals"],
                 "days_with_shortage": cov["days_with_shortage"], "changes": 0, "pending_leaves_not_applied": 0,
                 "warnings": ([f"{len(missing)} active employee(s) are not in the imported roster"] if missing else [])
                 + ([f"Shortage on {cov['days_with_shortage']} day(s) - see coverage report"] if cov["days_with_shortage"] else [])}
    audit.log(db, "roster_imported", "roster_version", v.id,
              {"roster_key": key, "version_no": version_no, "file": filename, "cells": len(cells),
               "shortage": cov["totals"]}, version_id=v.id, actor=actor)
    info = {"version_id": v.id, "start_date": str(start), "end_date": str(end), "cells": len(cells),
            "employees": len({e for e, _ in cells}), "unknown_employee_ids": unknown[:MAX_REPORTED],
            "not_in_file_active_employees": len(missing), "shortage_total": cov["totals"]["total"],
            "days_with_shortage": cov["days_with_shortage"]}
    return v, info


def import_workbook_bytes(db: Session, data: bytes, *, actor: str, dry_run: bool = True, week_start: date | None = None,
                          employees: bool = True, requirements: bool = True, roster: bool = True,
                          filename: str | None = None) -> dict:
    if not data:
        raise RosterError("The file is empty", 422)
    if len(data) > MAX_BYTES:
        raise RosterError(f"File is too large (max {MAX_BYTES // (1024 * 1024)} MB)", 413)
    if data[:2] != b"PK":
        raise RosterError("Please upload an Excel .xlsx workbook", 415)
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise RosterError(f"Could not read the workbook: {exc}", 422) from exc

    out: dict = {"dry_run": dry_run, "sheets": list(wb.sheetnames), "employees": None, "requirements": None,
                 "leaves": None, "roster": None, "errors": [], "warnings": []}
    try:
        if employees and any(n in wb.sheetnames for n in SHEET_NAMES):
            created, updated, skipped, errs = import_workbook(io.BytesIO(data), db=db, commit=False)
            out["employees"] = {"created": created, "updated": updated, "skipped": skipped}
            out["errors"] += errs
        if requirements and ("Daily_Requirement" in wb.sheetnames or "Leave_Applications" in wb.sheetnames):
            res = import_inputs(io.BytesIO(data), db=db, commit=False)
            out["requirements"] = res["requirements"]
            out["leaves"] = {"created": res["leaves_created"], "updated": res["leaves_updated"]}
            out["errors"] += res["errors"]
        if roster and any(n in wb.sheetnames for n in GRID_SHEETS):
            grid_errors: list[str] = []
            cells = parse_grid(wb, week_start, grid_errors)
            out["errors"] += grid_errors
            if cells:
                _, info = _create_draft(db, cells, actor=actor, reason=None, filename=filename, warnings=out["warnings"])
                out["roster"] = info
            else:
                out["warnings"].append("No roster rows were imported" + ("" if week_start else " - pick the roster start date"))
        if all(out[k] is None for k in ("employees", "requirements", "leaves", "roster")):
            why = "; ".join(out["errors"][:3] + out["warnings"][:1])
            raise RosterError("Nothing to import" + (f": {why}" if why else
                              ": the workbook has none of Employee_Master, Daily_Requirement, Leave_Applications "
                              "or Weekly_Roster (or those parts were switched off)"), 422)
        if not dry_run:
            audit.log(db, "workbook_imported", "import", filename or "workbook",
                      {k: out[k] for k in ("employees", "requirements", "leaves")} | {"errors": len(out["errors"])},
                      actor=actor)
            db.commit()
        else:
            db.rollback()
    except Exception:
        db.rollback()
        raise
    finally:
        wb.close()
    out["errors"] = out["errors"][:MAX_REPORTED]
    out["warnings"] = out["warnings"][:MAX_REPORTED]
    if dry_run and out["roster"]:
        out["roster"] = {**out["roster"], "version_id": None}  # nothing was saved
    return out
