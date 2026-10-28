"""Tests that run against the 5 demo roster workbooks."""
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models.employee import Employee, RolePool
from app.scripts.import_employees import import_workbook

EXPECTED_SHEETS = [
    "README", "Employee_Master", "Daily_Requirement", "Leave_Applications",
    "Roster_Draft", "Coverage_Report", "Fairness_Summary",
]


def by_key(files, key):
    return next(f for f in files if key in f.name)


def test_five_demo_workbooks_exist(demo_files):
    assert len(demo_files) == 5


def test_all_workbooks_have_expected_sheets_and_unique_ids(demo_files):
    for f in demo_files:
        wb = load_workbook(f, data_only=True)
        assert wb.sheetnames == EXPECTED_SHEETS, f.name
        ids = [r[0] for r in wb["Employee_Master"].iter_rows(min_row=2, values_only=True)]
        assert len(ids) >= 200, f.name
        assert len(ids) == len(set(ids)), f"duplicate employee IDs in {f.name}"
        assert wb["Daily_Requirement"].max_row - 1 == 14


def test_night_target_is_14(demo_files):
    for f in demo_files:
        wb = load_workbook(f, data_only=True)
        assert {r[4] for r in wb["Daily_Requirement"].iter_rows(min_row=2, values_only=True)} == {14}


def test_roster_has_no_night_to_morning_and_max_3_consecutive_nights(demo_files):
    for f in demo_files:
        ws = load_workbook(f, data_only=True)["Roster_Draft"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            codes = row[3:]
            streak = 0
            for prev, cur in zip(codes, codes[1:]):
                assert not (prev == "N" and str(cur).startswith("M")), f"{f.name} {row[0]}"
            for c in codes:
                streak = streak + 1 if c == "N" else 0
                assert streak <= 3, f"{f.name} {row[0]} >3 nights in a row"


def test_roster_respects_weekly_off_leave_and_eligibility(demo_files):
    wd = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    for f in demo_files:
        wb = load_workbook(f, data_only=True)
        emps = {r[0]: r for r in wb["Employee_Master"].iter_rows(min_row=2, values_only=True)}
        days = [r[0] for r in wb["Daily_Requirement"].iter_rows(min_row=2, values_only=True)]
        for row in wb["Roster_Draft"].iter_rows(min_row=2, values_only=True):
            e = emps[row[0]]
            for d, code in zip(days, row[3:]):
                code = str(code)
                not_employed = e[12] == "No" or d.date() < e[8].date()
                if not_employed:
                    assert code == "-", f"{f.name} {row[0]} should be '-' on {d.date()}"
                    continue
                if wd[d.weekday()] == e[6]:
                    assert code == "WO"
                if code.startswith("N"):
                    assert e[11] == "Yes", f"{f.name}: non-night-eligible {row[0]} on night"
        # approved full-day leave must show as leave code in the roster
        roster = {r[0]: r[3:] for r in wb["Roster_Draft"].iter_rows(min_row=2, values_only=True)}
        code_for = {"Sick": "SL", "Casual": "CL", "Annual": "AL"}
        for lv in wb["Leave_Applications"].iter_rows(min_row=2, values_only=True):
            _, emp, kind, start, end, half, _, status, _ = lv
            if status != "Approved" or half == "Yes":
                continue
            for i, d in enumerate(days):
                if start.date() <= d.date() <= end.date():
                    e = emps[emp]
                    if wd[d.weekday()] != e[6] and d.date() >= e[8].date():
                        assert roster[emp][i] == code_for[kind]


def test_scenario_expectations(demo_files):
    def totals(key):
        ws = load_workbook(by_key(demo_files, key), data_only=True)["Coverage_Report"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        return {"night": sum(r[10] for r in rows), "senior": sum(r[14] for r in rows),
                "short_days": sum(1 for r in rows if r[-1] == "SHORTAGE")}

    assert totals("01_standard") == {"night": 0, "senior": 0, "short_days": 0}
    assert totals("02_night_shortage")["night"] > 0
    assert totals("03_senior_tight")["senior"] > 0


def test_standard_night_fairness_spread(demo_files):
    ws = load_workbook(by_key(demo_files, "01_standard"), data_only=True)["Fairness_Summary"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[2] == "Yes"]
    nights = [r[5] for r in rows]
    assert max(nights) - min(nights) <= 4


def test_leave_heavy_has_many_applications_with_mixed_types(demo_files):
    ws = load_workbook(by_key(demo_files, "04_leave_heavy"), data_only=True)["Leave_Applications"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 45
    assert {r[2] for r in rows} == {"Sick", "Casual", "Annual"}
    assert {r[7] for r in rows} == {"Approved", "Pending"}
    assert any(r[5] == "Yes" for r in rows)  # at least one half-day leave


def test_no_overlapping_leave_per_employee(demo_files):
    for f in demo_files:
        spans = {}
        for r in load_workbook(f, data_only=True)["Leave_Applications"].iter_rows(min_row=2, values_only=True):
            spans.setdefault(r[1], []).append((r[3], r[4]))
        for emp, lst in spans.items():
            lst.sort()
            for (a1, b1), (a2, b2) in zip(lst, lst[1:]):
                assert a2 > b1, f"{f.name}: overlapping leave for {emp}"


def test_leave_codes_appear_in_leave_heavy_roster(demo_files):
    ws = load_workbook(by_key(demo_files, "04_leave_heavy"), data_only=True)["Roster_Draft"]
    codes = {c for r in ws.iter_rows(min_row=2, values_only=True) for c in r[3:]}
    assert {"SL", "CL", "AL"} <= codes


# ------------------------------------------------------------ importer ---

@pytest.mark.parametrize("key", ["01_standard", "02_night_shortage", "03_senior_tight",
                                 "04_leave_heavy", "05_new_joiners"])
def test_import_each_workbook_into_sqlite(demo_files, sqlite_session, key):
    path = by_key(demo_files, key)
    expected = load_workbook(path, data_only=True)["Employee_Master"].max_row - 1
    created, updated, skipped, errors = import_workbook(path, db=sqlite_session)
    assert (created, updated, skipped, errors) == (expected, 0, 0, [])
    assert sqlite_session.scalar(select(func.count()).select_from(Employee)) == expected


def test_import_is_idempotent(demo_files, sqlite_session):
    path = by_key(demo_files, "01_standard")
    import_workbook(path, db=sqlite_session)
    created, updated, _, errors = import_workbook(path, db=sqlite_session)
    assert (created, updated, errors) == (0, 200, [])


def test_import_maps_roles_eligibility_and_dates(demo_files, sqlite_session):
    import_workbook(by_key(demo_files, "01_standard"), db=sqlite_session)
    seniors = sqlite_session.scalar(
        select(func.count()).select_from(Employee).where(Employee.role_pool == RolePool.senior_support))
    nights = sqlite_session.scalar(
        select(func.count()).select_from(Employee).where(Employee.is_shift_eligible_night.is_(True)))
    assert seniors == 24 and nights == 90
    emp = sqlite_session.scalar(select(Employee).where(Employee.employee_id == "EMP001"))
    assert isinstance(emp.joining_date, date) and emp.weekly_off_day is not None


def test_import_new_joiners_and_inactive(demo_files, sqlite_session):
    import_workbook(by_key(demo_files, "05_new_joiners"), db=sqlite_session)
    inactive = sqlite_session.scalar(
        select(func.count()).select_from(Employee).where(Employee.is_active.is_(False)))
    future = sqlite_session.scalar(
        select(func.count()).select_from(Employee).where(Employee.joining_date > date(2026, 10, 26)))
    new_night = sqlite_session.scalar(
        select(func.count()).select_from(Employee)
        .where(Employee.notes == "New joiner - onboarding", Employee.is_shift_eligible_night.is_(True)))
    assert inactive == 6 and future > 0 and new_night == 0


def test_import_rejects_workbook_without_employee_master(tmp_path, sqlite_session):
    from openpyxl import Workbook
    p = tmp_path / "bad.xlsx"
    wb = Workbook()
    wb.active.title = "Other"
    wb.save(p)
    with pytest.raises(ValueError, match="Employee_Master"):
        import_workbook(p, db=sqlite_session)
