"""Step 10: final end-to-end run of the whole business flow through the real HTTP API (auth ON, approval ON)."""
import io
from datetime import timedelta

from openpyxl import load_workbook

from tests.helpers import START
from tests.test_leave_and_versions import Q, client  # noqa: F401  (fixture)
from tests.test_production import API, PW, hdr, sec  # noqa: F401  (fixtures)


def test_full_lifecycle(sec):
    # 1. admin creates a manager and an employee login
    admin = hdr(sec, "admin")
    assert sec.post(f"{API}/auth/users", headers=admin, json={"username": "ops_lead", "password": PW, "role": "manager"}).status_code == 201
    assert sec.post(f"{API}/auth/users", headers=admin, json={"username": "rahim", "password": PW, "role": "employee",
                                                              "employee_id": "T010"}).status_code == 201
    mgr, rahim = hdr(sec, "ops_lead"), hdr(sec, "rahim")
    assert sec.get(f"{API}/auth/me", headers=rahim).json()["employee_id"] == "T010"

    # 2. manager hires someone new, generates, reviews and publishes the roster
    assert sec.post(f"{API}/employees?auto_roster=false", headers=mgr,
                    json={"employee_id": "T900", "full_name": "New Hire", "joining_date": "2026-11-02"}).status_code == 201
    r = sec.post(f"{API}/roster/generate", headers=mgr, json={"start_date": str(START), "days": 14, **Q})
    assert r.status_code == 201
    v = r.json()
    assert sec.get(f"{API}/roster/versions", headers=rahim).json() == []  # drafts stay private
    assert sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr, json={"acknowledge_shortage": True}).status_code == 409  # approve first
    ack = {"acknowledge_shortage": True}
    assert sec.post(f"{API}/roster/versions/{v['id']}/approve", headers=mgr, json={**ack, "note": "ok"}).status_code == 200
    pub = sec.post(f"{API}/roster/versions/{v['id']}/publish", headers=mgr, json=ack)
    assert pub.status_code == 200 and pub.json()["status"] == "published" and pub.json()["published_by"] == "ops_lead"
    assert [x["id"] for x in sec.get(f"{API}/roster/versions", headers=rahim).json()] == [v["id"]]
    assert sec.put(f"{API}/roster/versions/{v['id']}/assignments", headers=mgr,
                   json={"employee_id": "T010", "work_date": str(START), "code": "WO", "reason": "attempted edit"}).status_code == 409  # immutable

    # 3. employee applies for leave in the published window, manager approves -> new draft version
    lv = sec.post(f"{API}/leaves", headers=rahim, json={"employee_id": "T010", "leave_type": "Casual",
                                                        "start_date": str(START + timedelta(days=7)),
                                                        "end_date": str(START + timedelta(days=8))})
    assert lv.status_code == 201
    lid = lv.json()["application_id"]
    ap = sec.post(f"{API}/leaves/{lid}/approve", headers=mgr, json={})
    assert ap.status_code == 200 and ap.json()["leave"]["status"] == "Approved"
    new_id = ap.json()["roster_version"]["id"] if isinstance(ap.json().get("roster_version"), dict) else ap.json().get("roster_version")
    assert new_id and new_id != v["id"]
    cells = {r_["employee_id"]: r_["codes"] for r_ in sec.get(f"{API}/roster/versions/{new_id}", headers=mgr).json()["grid"]["rows"]}
    assert cells["T010"][7] == "CL" and cells["T010"][8] == "CL"
    assert sec.get(f"{API}/roster/versions/{v['id']}", headers=rahim).json()["version"]["status"] in ("published", "superseded")

    # 4. exports (manager) and the employee's own weekly sheet
    for path, sheet in (("employees.xlsx", "Employees"), ("leaves.xlsx", "Leave Applications"),
                        (f"coverage.xlsx?version_id={v['id']}", "Shift Coverage")):
        r = sec.get(f"{API}/exports/{path}", headers=mgr)
        assert r.status_code == 200 and sheet in load_workbook(io.BytesIO(r.content)).sheetnames, path
    r = sec.get(f"{API}/exports/roster/weekly.xlsx?week_start={START}&version_id={v['id']}", headers=rahim)
    assert r.status_code == 200 and "T010" in [c.value for c in load_workbook(io.BytesIO(r.content))["Weekly Roster"]["A"]]

    # 5. audit trail names every actor
    log = sec.get(f"{API}/roster/audit-log?limit=500", headers=mgr).json()
    seen = {(a["action"], a["actor"]) for a in log}
    for expected in (("roster_generated", "ops_lead"), ("leave_applied", "rahim"), ("leave_approved", "ops_lead"),
                     ("report_exported", "ops_lead"), ("login", "rahim"), ("employee_created", "ops_lead")):
        assert any(a == expected[0] and who == expected[1] for a, who in seen), f"missing audit {expected}: {sorted(seen)[:20]}"

    # 6. disabling the account takes effect immediately, even with a still-valid token
    assert sec.patch(f"{API}/auth/users/rahim", headers=admin, json={"is_active": False}).status_code == 200
    assert sec.get(f"{API}/auth/me", headers=rahim).status_code == 401
    assert sec.post(f"{API}/auth/login", json={"username": "rahim", "password": PW}).status_code == 401
